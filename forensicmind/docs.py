"""Generación de documentos: dictamen en Word (.docx) y ficha vertical 9:16 (.png) con Pillow."""
from __future__ import annotations

import io
import re
import textwrap
from datetime import datetime

from . import config
from .security import quitar_markdown

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # pragma: no cover
    Document = None

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = None

_RE_NO_TEXTO = re.compile("[\U00010000-\U0010ffff☀-➿‍️⭐⬆⏰⏱]")


def limpiar_para_imagen(texto: str) -> str:
    """Quita markdown y emojis (las fuentes TTF básicas no los dibujan)."""
    t = quitar_markdown(texto)
    t = _RE_NO_TEXTO.sub("", t)
    return re.sub(r"\s+", " ", t).strip()


# ------------------------------------------------------------------------------ Word
def _bullets(doc, items):
    for it in items:
        doc.add_paragraph(quitar_markdown(it), style="List Bullet")


def generar_docx(
    folio: str, relato: str, triage: dict, guias_por_modulo: list[dict], resultado_lft=None,
    oferta: float = 0.0, fundamentos: list[dict] | None = None, hash_ev: str = "",
) -> bytes:
    if Document is None:
        return b""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("ForensicMind & LexGuía — Manual de triage paso a paso", 0)
    p = doc.add_paragraph()
    p.add_run("Aviso: ").bold = True
    p.add_run(config.AVISO_LEGAL)
    doc.add_paragraph(f"Expediente: {folio}  |  Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"Autoría del sistema: {config.AUTOR_SISTEMA}")
    if triage.get("categorias_ordenadas"):
        doc.add_paragraph("Materias detectadas (por urgencia): " + "; ".join(triage["categorias_ordenadas"]))
    if hash_ev:
        doc.add_paragraph(f"Sello SHA-256 del relato: {hash_ev}")

    if relato.strip():
        doc.add_heading("Tu relato", level=1)
        doc.add_paragraph(relato.strip())

    if resultado_lft is not None:
        r = resultado_lft
        doc.add_heading("Estimación de derechos laborales (LFT)", level=1)
        doc.add_paragraph(
            f"Salario diario: ${r.salario_diario:,.2f} | Salario diario integrado: ${r.sdi:,.2f} | Antigüedad: {r.antiguedad:.2f} años"
        )
        tabla = doc.add_table(rows=1, cols=4)
        tabla.style = "Light Grid Accent 1"
        for i, h in enumerate(("Concepto", "Fundamento", "Monto (MXN)", "Cálculo")):
            tabla.rows[0].cells[i].text = h
        for c in r.conceptos:
            fila = tabla.add_row().cells
            fila[0].text, fila[1].text = c["concepto"], c["fundamento"]
            fila[2].text, fila[3].text = f"${c['monto']:,.2f}", c.get("nota", "")
        doc.add_paragraph()
        doc.add_paragraph(f"Total estimado por ley (finiquito + liquidación): ${r.total_liquidacion:,.2f} MXN")
        if r.es_despido:
            doc.add_paragraph(f"Con salarios vencidos hasta 12 meses (solo si hay juicio y se gana): ${r.total_con_salarios_vencidos:,.2f} MXN")
        if oferta and oferta < r.total_liquidacion:
            doc.add_paragraph(f"Oferta del patrón: ${oferta:,.2f} MXN → diferencia frente a la estimación legal: ${r.total_liquidacion - oferta:,.2f} MXN")
        doc.add_paragraph(
            f"Pagos por separación: ${r.isr_pago_separacion:,.2f}; exento de ISR estimado: ${r.isr_tope_exento:,.2f}; "
            f"gravable estimado: ${r.isr_gravable_estimado:,.2f} (informativo)."
        )
        _bullets(doc, r.notas)

    for idx, item in enumerate(guias_por_modulo, 1):
        g = item["guia"]
        doc.add_heading(f"Módulo {idx}: {quitar_markdown(item['titulo_prioridad'])}", level=1)
        doc.add_paragraph(quitar_markdown(g["radiografia"]))
        doc.add_heading("Explicación en español sencillo", level=2)
        _bullets(doc, g["explicacion_sencilla"])
        doc.add_heading("Qué hacer hoy mismo", level=2)
        _bullets(doc, g["fase_1_hoy"])
        doc.add_heading("Qué hacer esta semana", level=2)
        _bullets(doc, g["fase_2_semana"])
        doc.add_heading("Ante la autoridad", level=2)
        _bullets(doc, g["fase_3_autoridad"])
        doc.add_heading("Guion sugerido", level=2)
        doc.add_paragraph(quitar_markdown(g["guion"]))
        doc.add_heading("Errores que debes evitar", level=2)
        _bullets(doc, g["errores_fatales"])
        doc.add_paragraph(f"Dónde acudir: {g['autoridad']}")
        doc.add_paragraph(f"Especialista: {g['abogado_req']}")

    if fundamentos:
        doc.add_heading("Fundamento legal de referencia", level=1)
        for f in fundamentos:
            doc.add_paragraph(f"{f['ley']} Art. {f['articulo']} — {f['tema']}: {f['resumen']}", style="List Bullet")

    cont = config.leyes()["contactos"]
    doc.add_heading("Contactos útiles (verifica en sitios oficiales)", level=1)
    for k, etiqueta in (("emergencias", "Emergencias"), ("denuncia_anonima", "Denuncia anónima"), ("profedet", "PROFEDET"),
                        ("condusef", "CONDUSEF"), ("cndh", "CNDH")):
        doc.add_paragraph(f"{etiqueta}: {cont[k]}", style="List Bullet")
    doc.add_paragraph(f"Valores legales usados: salario mínimo ${config.leyes()['salario_minimo']['general']}, "
                      f"UMA ${config.leyes()['uma']['diaria']} (verificados el {config.leyes()['salario_minimo']['verificado']}).")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------ PNG 9:16
_RUTAS_FUENTE = {
    False: ["DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "Arial.ttf"],
    True: ["DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf",
           "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf", "Arialbd.ttf"],
}


def _fuente(tam: int, negrita: bool = False):
    for ruta in _RUTAS_FUENTE[negrita]:
        try:
            return ImageFont.truetype(ruta, tam)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=tam)
    except TypeError:  # Pillow antiguo
        return ImageFont.load_default()


def _envolver(draw, texto: str, fuente, ancho_max: int) -> list[str]:
    lineas, actual = [], ""
    for palabra in texto.split():
        prueba = f"{actual} {palabra}".strip()
        if draw.textlength(prueba, font=fuente) <= ancho_max:
            actual = prueba
        else:
            if actual:
                lineas.append(actual)
            actual = palabra
    if actual:
        lineas.append(actual)
    return lineas


def generar_ficha_png(folio: str, triage: dict, guia_principal: dict, resultado_lft=None, oferta: float = 0.0, hash_ev: str = "") -> bytes:
    """Ficha vertical 1080×1920 con el plan de acción. Devuelve bytes PNG."""
    if Image is None:
        return b""
    W, H, M = 1080, 1920, 60
    img = Image.new("RGB", (W, H), (10, 15, 24))
    d = ImageDraw.Draw(img)
    dorado, cian, blanco, gris, verde = (245, 175, 40), (110, 215, 250), (245, 245, 245), (190, 200, 215), (110, 235, 170)

    d.rectangle([20, 20, W - 20, H - 20], outline=dorado, width=4)
    d.rectangle([40, 40, W - 40, 235], fill=(26, 38, 60))
    d.text((M, 70), "ForensicMind & LexGuia", font=_fuente(58, True), fill=cian)
    d.text((M, 150), limpiar_para_imagen(config.AUTOR_SISTEMA), font=_fuente(28), fill=blanco)

    y = 270
    d.text((M, y), limpiar_para_imagen(f"Expediente {folio}  -  {datetime.now().strftime('%d/%m/%Y')}"), font=_fuente(32, True), fill=blanco)
    y += 55
    materias = triage.get("categorias_ordenadas") or ["Por determinar"]
    for i, mat in enumerate(materias[:6], 1):
        d.text((M, y), limpiar_para_imagen(f"{i}. {mat}"), font=_fuente(30), fill=dorado)
        y += 42

    y += 15
    d.rounded_rectangle([M, y, W - M, y + 170], radius=18, fill=(20, 50, 40))
    if resultado_lft is not None:
        d.text((M + 25, y + 22), limpiar_para_imagen(f"Estimado por ley (LFT): ${resultado_lft.total_liquidacion:,.2f} MXN"), font=_fuente(34, True), fill=blanco)
        if oferta and oferta < resultado_lft.total_liquidacion:
            d.text((M + 25, y + 78), limpiar_para_imagen(f"Te ofrecen: ${oferta:,.2f}  |  Diferencia: ${resultado_lft.total_liquidacion - oferta:,.2f}"), font=_fuente(30), fill=verde)
        d.text((M + 25, y + 125), "Estimacion informativa: verifica con la PROFEDET", font=_fuente(24), fill=gris)
    else:
        d.text((M + 25, y + 30), "Pagos a policias o extorsionadores: $0", font=_fuente(34, True), fill=blanco)
        d.text((M + 25, y + 90), "Via de defensa: denuncia, queja y reparacion del dano", font=_fuente(28), fill=verde)
    y += 210

    d.text((M, y), "PLAN DE ACCION", font=_fuente(40, True), fill=cian)
    y += 62
    f_txt = _fuente(30)
    limite_y = 1560
    pasos = guia_principal.get("fase_1_hoy", [])[:4] + guia_principal.get("fase_2_semana", [])[:2]
    for paso in pasos:
        lineas = _envolver(d, "• " + re.sub(r"^\d+\.\s*", "", limpiar_para_imagen(paso)), f_txt, W - 2 * M)
        if y + 38 * min(len(lineas), 4) > limite_y:
            d.text((M, y), "... (ver el manual completo en Word)", font=f_txt, fill=gris)
            y += 40
            break
        for i, lin in enumerate(lineas[:4]):
            if i == 3 and len(lineas) > 4:
                lin = lin.rstrip(".") + "..."
            d.text((M, y), lin, font=f_txt, fill=blanco)
            y += 38
        y += 14

    y = max(y + 10, 1590)
    d.text((M, y), "DONDE ACUDIR", font=_fuente(34, True), fill=cian)
    y += 48
    for lin in _envolver(d, limpiar_para_imagen(guia_principal.get("autoridad", "")), _fuente(28), W - 2 * M)[:3]:
        d.text((M, y), lin, font=_fuente(28), fill=dorado)
        y += 36

    cont = config.leyes()["contactos"]
    d.text((M, 1760), limpiar_para_imagen(f"Emergencias {cont['emergencias']}  |  Denuncia anonima {cont['denuncia_anonima']}  |  PROFEDET {cont['profedet']}"),
           font=_fuente(25), fill=verde)
    d.text((M, 1805), "Orientacion informativa; no sustituye a un abogado.", font=_fuente(24), fill=gris)
    if hash_ev:
        d.text((M, 1850), f"SHA-256: {hash_ev[:48]}", font=_fuente(22), fill=gris)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()
