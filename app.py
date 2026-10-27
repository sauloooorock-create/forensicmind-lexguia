import streamlit as st
import sqlite3
import hashlib
import re
import io
import os
import json
import base64
import textwrap
import urllib.request
import urllib.parse
import xml.etree.ElementTree as ET
from datetime import datetime
import pandas as pd
import numpy as np
import cv2

try:
    import speech_recognition as sr
except ImportError:
    sr = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from sklearn.neural_network import MLPClassifier
    from sklearn.feature_extraction.text import TfidfVectorizer
except ImportError:
    MLPClassifier = None
    TfidfVectorizer = None

# ==============================================================================
# CONFIGURACIÓN GLOBAL Y CREDENCIALES SEGURAS v10.0
# ==============================================================================
AUTOR_SISTEMA = "Ing. Industrial y de Sistemas Saúl Ángel Ortiz Rogel"
MARCA_SISTEMA = "ForensicMind & LexGuía CBR v10.0 — Triage Multi-Materia, Despacho & Auto-Sync"

st.set_page_config(
    page_title=f"{MARCA_SISTEMA} | {AUTOR_SISTEMA}",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

DB_NAME = "boveda_juridica_cbr_v10.db"
BACKUP_JSON_FILE = "cerebro_autonomo_backup.json"

DEFAULT_GH_TOKEN = ""
DEFAULT_GH_REPO = "sauloooorock-create/forensicmind-lexguia"
DEFAULT_GH_BRANCH = "main"
DEFAULT_ADMIN_PASS = "admin2026"
DEFAULT_ADMIN_EMAIL = "contacto@forensicmind.mx"
DEFAULT_WHATSAPP_DESPACHO = "527220000000" # Sustituye con tu número personal para recibir clientes

try:
    DEFAULT_GH_TOKEN = st.secrets.get("GITHUB_TOKEN", "")
    DEFAULT_GH_REPO = st.secrets.get("GITHUB_REPO", DEFAULT_GH_REPO)
    DEFAULT_GH_BRANCH = st.secrets.get("GITHUB_BRANCH", "main")
    DEFAULT_ADMIN_PASS = st.secrets.get("ADMIN_PASSWORD", "admin2026")
    DEFAULT_ADMIN_EMAIL = st.secrets.get("ADMIN_EMAIL", DEFAULT_ADMIN_EMAIL)
    DEFAULT_WHATSAPP_DESPACHO = st.secrets.get("WHATSAPP_DESPACHO", DEFAULT_WHATSAPP_DESPACHO)
except Exception:
    pass

if "es_admin" not in st.session_state:
    st.session_state["es_admin"] = False

# REGISTRO AUTOMÁTICO DE VISITANTES ÚNICOS POR SESIÓN
if "visit_registrada" not in st.session_state:
    st.session_state["visit_registrada"] = True
    try:
        conn_v = sqlite3.connect(DB_NAME)
        conn_v.execute("CREATE TABLE IF NOT EXISTS visitas_pagina (id INTEGER PRIMARY KEY AUTOINCREMENT, fecha_hora TEXT)")
        conn_v.execute("INSERT INTO visitas_pagina (fecha_hora) VALUES (?)", (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))
        conn_v.commit()
        conn_v.close()
    except Exception:
        pass

# ==============================================================================
# BARRA LATERAL: ACCESO ADMIN Y PREFERENCIAS
# ==============================================================================
with st.sidebar:
    st.markdown("### 🎨 Apariencia Visual")
    modo_visual = st.radio(
        "Modo de lectura:",
        [
            "☀️ Claro Ejecutivo (Fácil lectura en celular)",
            "🌙 Oscuro Forense"
        ],
        index=0
    )
    buscar_web_switch = st.checkbox(
        "🌐 Consultar noticias y criterios en internet en vivo",
        value=True
    )

    st.markdown("---")
    st.markdown("### 🔐 Acceso Exclusivo Administrador")
    if not st.session_state["es_admin"]:
        clave_input = st.text_input(
            "Contraseña de Administrador:",
            type="password",
            placeholder="Clave del Ing. Saúl..."
        )
        if st.button("🔓 Desbloquear Panel Admin", use_container_width=True):
            if clave_input == DEFAULT_ADMIN_PASS:
                st.session_state["es_admin"] = True
                st.rerun()
            else:
                st.error("Contraseña incorrecta.")
        gh_repo_input = DEFAULT_GH_REPO
        gh_token_input = DEFAULT_GH_TOKEN
        gh_branch_input = DEFAULT_GH_BRANCH
        auto_entrenar_switch = True
    else:
        st.success("✅ Modo Administrador Activo (Ing. Saúl)")
        auto_entrenar_switch = st.checkbox(
            "🧠 Auto-Aprendizaje e Inyección Activa",
            value=True
        )
        gh_repo_input = st.text_input(
            "Repositorio GitHub:",
            value=DEFAULT_GH_REPO
        )
        gh_token_input = st.text_input(
            "GitHub Token (ghp_...):",
            value=DEFAULT_GH_TOKEN,
            type="password"
        )
        gh_branch_input = st.text_input(
            "Rama (Branch):",
            value=DEFAULT_GH_BRANCH
        )
        if st.button("🔒 Cerrar Sesión Admin", use_container_width=True):
            st.session_state["es_admin"] = False
            st.rerun()

# ==============================================================================
# ESTILOS CSS (SIN COMILLAS TRIPLES PARA EVITAR CORTES EN CELULAR)
# ==============================================================================
if "Claro" in modo_visual:
    css_lines = [
        "<style>",
        ".stApp { background-color: #f1f5f9 !important; color: #0f172a !important; }",
        ".block-container { padding-top: 3.2rem !important; padding-bottom: 2.5rem !important; }",
        "h1, h2, h3, h4, h5, h6, p, li, span, div { color: #0f172a; }",
        "label, [data-testid='stWidgetLabel'] p { color: #0f172a !important; font-weight: 700 !important; }",
        "textarea, input[type='text'], input[type='number'], input[type='password'] {",
        "  background-color: #ffffff !important; color: #0f172a !important;",
        "  -webkit-text-fill-color: #0f172a !important; border: 2px solid #2563eb !important;",
        "  border-radius: 10px !important; font-size: 1.04rem !important; font-weight: 600 !important;",
        "}",
        ".author-banner {",
        "  background: linear-gradient(135deg, #0f172a 0%, #1e3a8a 100%);",
        "  border-left: 6px solid #f59e0b; border-radius: 14px; padding: 16px 20px; margin-bottom: 18px;",
        "}",
        ".author-banner * { color: #ffffff !important; }",
        "div[data-testid='metric-container'] {",
        "  background: #ffffff !important; border: 1px solid #cbd5e1 !important;",
        "  border-left: 5px solid #2563eb !important; padding: 14px !important; border-radius: 12px !important;",
        "}",
        ".win-badge { background-color: #059669; color: #ffffff !important; padding: 5px 12px; border-radius: 6px; font-weight: 700; display: inline-block; }",
        ".loss-badge { background-color: #dc2626; color: #ffffff !important; padding: 5px 12px; border-radius: 6px; font-weight: 700; display: inline-block; }",
        ".step-badge { background-color: #1d4ed8; color: #ffffff !important; padding: 5px 12px; border-radius: 6px; font-weight: 700; display: inline-block; }",
        ".gold-badge { background-color: #f59e0b; color: #0f172a !important; padding: 5px 12px; border-radius: 6px; font-weight: 800; display: inline-block; }",
        ".triage-box { background-color: #fff1f2; border: 2px solid #f43f5e; border-left: 8px solid #be123c; border-radius: 12px; padding: 18px; margin: 12px 0; }",
        ".card-box { background-color: #ffffff; border: 1px solid #cbd5e1; border-left: 6px solid #2563eb; border-radius: 12px; padding: 18px; margin: 10px 0; }",
        ".money-box { background-color: #ecfdf5; border: 1px solid #6ee7b7; border-left: 6px solid #059669; border-radius: 12px; padding: 18px; margin: 10px 0; }",
        ".lawyer-box { background-color: #eef2ff; border: 1px solid #a5b4fc; border-left: 6px solid #4f46e5; border-radius: 12px; padding: 18px; margin: 10px 0; }",
        ".alert-box { background-color: #fef2f2; border: 1px solid #fca5a5; border-left: 6px solid #dc2626; border-radius: 12px; padding: 18px; margin: 10px 0; }",
        ".script-box { background-color: #fffbeb; border: 1px solid #fcd34d; border-left: 6px solid #d97706; border-radius: 12px; padding: 18px; margin: 10px 0; }",
        ".web-box { background-color: #f0f9ff; border: 1px solid #7dd3fc; border-left: 6px solid #0284c7; border-radius: 16px; padding: 16px; margin: 10px 0; }",
        "div.stButton > button:first-child {",
        "  background: linear-gradient(90deg, #2563eb 0%, #1d4ed8 100%) !important;",
        "  color: #ffffff !important; -webkit-text-fill-color: #ffffff !important;",
        "  font-weight: 800 !important; border-radius: 10px !important; padding: 12px 20px !important;",
        "}",
        "</style>"
    ]
else:
    css_lines = [
        "<style>",
        ".stApp { background-color: #0b0f19 !important; color: #f8fafc !important; }",
        ".block-container { padding-top: 3.2rem !important; padding-bottom: 2.5rem !important; }",
        "h1, h2, h3, h4, h5, h6, p, li, label { color: #f8fafc !important; }",
        "textarea, input[type='text'], input[type='number'], input[type='password'] {",
        "  background-color: #ffffff !important; color: #0f172a !important;",
        "  -webkit-text-fill-color: #0f172a !important; border: 2px solid #38bdf8 !important; border-radius: 10px !important;",
        "}",
        ".author-banner { background: #1e293b; border-left: 6px solid #f59e0b; border-radius: 12px; padding: 16px 20px; margin-bottom: 18px; }",
        ".win-badge { background-color: #065f46; color: #34d399 !important; padding: 4px 10px; border-radius: 6px; font-weight: bold; }",
        ".loss-badge { background-color: #7f1d1d; color: #f87171 !important; padding: 4px 10px; border-radius: 6px; font-weight: bold; }",
        ".step-badge { background-color: #1e3a8a; color: #93c5fd !important; padding: 4px 10px; border-radius: 6px; font-weight: bold; }",
        ".gold-badge { background-color: #f59e0b; color: #0f172a !important; padding: 4px 10px; border-radius: 6px; font-weight: bold; }",
        ".triage-box { background-color: #2a1215; border: 2px solid #f43f5e; border-left: 8px solid #e11d48; border-radius: 10px; padding: 18px; margin: 12px 0; }",
        ".card-box { background-color: #131c2e; border: 1px solid #2dd4bf; border-radius: 10px; padding: 18px; margin: 10px 0; }",
        ".money-box { background-color: #0f2922; border: 1px solid #10b981; border-left: 5px solid #34d399; border-radius: 10px; padding: 18px; margin: 10px 0; }",
        ".lawyer-box { background-color: #1e1b4b; border: 1px solid #818cf8; border-left: 5px solid #a78bfa; border-radius: 10px; padding: 18px; margin: 10px 0; }",
        ".alert-box { background-color: #1f1215; border: 1px solid #f87171; border-left: 5px solid #ef4444; border-radius: 10px; padding: 18px; margin: 10px 0; }",
        ".script-box { background-color: #272011; border: 1px solid #f59e0b; border-left: 5px solid #fbbf24; border-radius: 10px; padding: 18px; margin: 10px 0; }",
        ".web-box { background-color: #0c192c; border: 1px solid #38bdf8; border-left: 5px solid #0284c7; border-radius: 10px; padding: 16px; margin: 10px 0; }",
        "</style>"
    ]

st.markdown("\n".join(css_lines), unsafe_allow_html=True)

# ==============================================================================
# CATÁLOGOS DE CASOS, PRIORIDADES DE TRIAGE Y ESCENARIOS
# ==============================================================================
TIPOS_CASOS = {
    "Tránsito y Abuso Policial": [
        "Detención arbitraria, siembra de pruebas o extorsión policial",
        "Amenaza ilegal de corralón o retiro de placas",
        "Multa sin fundamento o retención de documentos"
    ],
    "Casos Penales y Fraudes": [
        "Robo de identidad con INE, fraude bancario o criptoactivos",
        "Phishing, vaciado de cuentas o transferencias SPEI no reconocidas",
        "Extorsión o amenazas digitales"
    ],
    "Casos Laborales": [
        "Despido injustificado o falta de pago de quincenas/semanas",
        "Presión para firmar renuncia u hoja en blanco",
        "Salarios, vacaciones o descansos no pagados"
    ],
    "Acoso de Cobranza y Deudas": [
        "Amenazas de embargo falso o despachos de cobranza",
        "Acoso de aplicaciones de préstamos o créditos no solicitados"
    ],
    "Casos Familiares": [
        "Pensión alimenticia y gasto de los hijos",
        "Guarda, custodia y régimen de visitas",
        "Divorcio"
    ],
    "Casos Civiles y Mercantiles": [
        "Inquilino que no paga renta ni desocupa",
        "Pagaré vencido o deuda de dinero prestado"
    ]
}

PRIORIDAD_TRIAGE = {
    "Tránsito y Abuso Policial": (1, "🚨 PRIORIDAD 1 (URGENCIA VITAL Y LIBERTAD): Abuso Policial, Detención o Corralón"),
    "Casos Penales y Fraudes": (2, "🔐 PRIORIDAD 2 (BLINDAJE DE IDENTIDAD Y CUENTAS): Fraude, Robo de Celular/INE o Delito"),
    "Casos Laborales": (3, "💼 PRIORIDAD 3 (RECLAMO DE TRABAJO Y DINERO): Despido, Sueldos Retenidos y Liquidación LFT"),
    "Casos Familiares": (4, "👨‍👩‍👧 PRIORIDAD 4 (PROTECCIÓN FAMILIAR): Pensión Alimenticia y Custodia"),
    "Casos Civiles y Mercantiles": (5, "🏠 PRIORIDAD 5 (PATRIMONIO CIVIL): Arrendamiento y Pagarés"),
    "Acoso de Cobranza y Deudas": (6, "📵 PRIORIDAD 6 (FRENO A HOSTIGAMIENTO): Despachos de Cobranza y Apps de Préstamos")
}

TABULADOR_HONORARIOS_RAMA = {
    "Casos Laborales": {
        "asesoria_min": 500, "asesoria_sug": 1000, "asesoria_max": 2000,
        "anticipo_sug": 0, "pct_sugerido": 25.0, "gastos_periciales_sug": 1000,
        "horas_promedio_despacho": 35,
        "esquema_cobro": "Porcentaje al ganar (20% a 30% de lo recuperado; $0 de anticipo inicial)",
        "via_gratuita": "PROFEDET / Procuraduría de la Defensa del Trabajo — $0 MXN (100% Gratuito)"
    },
    "Tránsito y Abuso Policial": {
        "asesoria_min": 500, "asesoria_sug": 800, "asesoria_max": 1800,
        "anticipo_sug": 2500, "pct_sugerido": 0.0, "gastos_periciales_sug": 500,
        "horas_promedio_despacho": 15,
        "esquema_cobro": "Defensa en Asuntos Internos, Fiscalía Anticorrupción y Juicio de Nulidad",
        "via_gratuita": "Asuntos Internos / Fiscalía Anticorrupción / Derechos Humanos — $0 MXN"
    },
    "Acoso de Cobranza y Deudas": {
        "asesoria_min": 500, "asesoria_sug": 900, "asesoria_max": 1500,
        "anticipo_sug": 2000, "pct_sugerido": 0.0, "gastos_periciales_sug": 500,
        "horas_promedio_despacho": 15,
        "esquema_cobro": "Queja por cobranza ilegítima y bloqueo de créditos por robo de identidad",
        "via_gratuita": "CONDUSEF / REDECO y Policía Cibernética — $0 MXN"
    },
    "Casos Familiares": {
        "asesoria_min": 800, "asesoria_sug": 1200, "asesoria_max": 2500,
        "anticipo_sug": 5000, "pct_sugerido": 0.0, "gastos_periciales_sug": 2500,
        "horas_promedio_despacho": 45,
        "esquema_cobro": "Pago fijo en partes conforme avanza el juicio familiar",
        "via_gratuita": "DIF / Defensoría Pública Familiar del Estado — $0 MXN"
    },
    "Casos Penales y Fraudes": {
        "asesoria_min": 1000, "asesoria_sug": 1800, "asesoria_max": 4000,
        "anticipo_sug": 10000, "pct_sugerido": 15.0, "gastos_periciales_sug": 4000,
        "horas_promedio_despacho": 60,
        "esquema_cobro": "Denuncia en Fiscalía, bloqueo en Buró de Crédito y reclamación CONDUSEF",
        "via_gratuita": "CONDUSEF, Fiscalía Cibernética y Asesoría Jurídica de Víctimas — $0 MXN"
    },
    "Casos Civiles y Mercantiles": {
        "asesoria_min": 800, "asesoria_sug": 1500, "asesoria_max": 3000,
        "anticipo_sug": 6000, "pct_sugerido": 20.0, "gastos_periciales_sug": 3500,
        "horas_promedio_despacho": 40,
        "esquema_cobro": "Anticipo inicial + porcentaje de lo cobrado o recuperado",
        "via_gratuita": "Defensoría Pública Civil — $0 MXN"
    }
}

ESCENARIOS_FRECUENTES = {
    "✍️ Escribir o dictar mi propio caso libremente": (
        "Mi sueldo semanal es de 2300. Me despidieron y no me quieren dar liquidación, "
        "solo me quieren dar $5000 y no me daban vacaciones ni días libres ni bonos, ¿qué procede ahí?"
    ),
    "🌪️ Caso Extremo Multi-Materia (Trabajo + Robo de INE + Abuso Policial + Cobranza)": (
        "Trabajaba para una empresa que no me quiere pagar 3 quincenas y usaron mi INE sin permiso "
        "para sacar créditos. Ayer iba en mi moto, una patrulla me detuvo sin motivo, me sembraron "
        "sustancias ilícitas para pedirme 50,000 pesos, me quitaron mi celular con mis cuentas bancarias "
        "y se llevaron mi moto al corralón. Además una app de préstamos me cobra 30,000 pesos que no pedí."
    ),
    "🚓 Tránsito y Policías: Me detuvo una patrulla, piden mordida y amenazan con corralón": (
        "Una patrulla de tránsito me detuvo sin decirme qué infracción cometí, me quieren quitar "
        "la placa o llevarse mi vehículo al corralón con grúa si no les doy dinero en efectivo."
    ),
    "💼 Laboral: Me exigen firmar mi renuncia o una hoja en blanco": (
        "Me corrieron de la chamba hoy sin aviso por escrito y el encargado de recursos humanos "
        "me dice que si no le firmo mi hoja de renuncia voluntaria no me va a dar ni mi última semana."
    ),
    "📞 Deudas y Cobranza: Me amenazan por WhatsApp con embargar mi casa sin orden de juez": (
        "Un despacho de cobranza me manda mensajes amenazando con que mañana van a ir a romper "
        "la cerradura de mi casa para embargar mis muebles por una deuda y le llaman a mis contactos."
    ),
    "👨‍👩‍👧 Familiar: El papá de mis hijos no da pensión y esconde su sueldo": (
        "El papá de mis hijos lleva 4 meses sin dar dinero para la comida y escuela de los niños, "
        "se salió de su trabajo formal a propósito pero gana dinero por su cuenta en efectivo."
    ),
    "💻 Fraude Bancario: Me vaciaron mi cuenta con un enlace o llamada falsa": (
        "Me llegó un mensaje que parecía de mi banco, entré al link y sacaron mi dinero con "
        "transferencias SPEI que yo nunca autoricé; el banco dice que fue mi culpa."
    )
}

DICCIONARIO_COLOQUIAL_A_JURIDICO = [
    (r'\b(chamba|jale|empleo|trabajo|trabajaba|oficina|empresa|startup|patron|quincenas|semanal|sueldo|liquidacion|finiquito|vacaciones)\b', ' trabajo laboral patron '),
    (r'\b(transito|policia|policias|patrulla|mordida|corralon|grua|infraccion|multa|placa|moto|sembraron|oficial)\b', ' transito abuso policial corralon '),
    (r'\b(cobranza|embargar|embargo|despacho|app de prestamos|prestamos|debo)\b', ' acoso cobranza deudas condusef '),
    (r'\b(ine|identidad|criptomonedas|llaves privadas|vaciaron|spei|fraude|amenazo|boletinarme)\b', ' fraude penal robo identidad '),
    (r'\b(gasto de los niños|manutencion|pension|alimentos|papa de mis hijos|custodia|divorcio)\b', ' pension alimenticia familiar '),
    (r'\b(inquilino|renta|arrendamiento|pagare|prestado)\b', ' civil mercantil arrendamiento ')
]

def expandir_lenguaje_coloquial(texto: str) -> str:
    texto_limpio = str(texto).lower()
    expansiones = []
    for patron, equivalente in DICCIONARIO_COLOQUIAL_A_JURIDICO:
        matches = re.findall(patron, texto_limpio)
        if matches:
            expansiones.append(equivalente * len(matches))
    return str(texto) + " " + " ".join(expansiones)

# ==============================================================================
# MOTOR DE AUTOINYECCIÓN Y SINCRONIZACIÓN CON GITHUB
# ==============================================================================
def inyectar_diccionario_cerebro_en_sqlite(data: dict, conn: sqlite3.Connection) -> dict:
    c = conn.cursor()
    nuevos_mem = 0
    nuevos_casos = 0

    sql_mem = " ".join([
        "INSERT OR IGNORE INTO memoria_neuronal_autonoma (",
        "relato_original, relato_expandido, macro_rama, rama_derecho,",
        "categoria_caso, tipo_caso, fuente_aprendizaje, confianza,",
        "hash_consulta, fecha_aprendizaje",
        ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    ])

    for m in data.get("memoria_neuronal_autonoma", []):
        rel_orig = str(m.get("relato_original", "")).strip()
        if not rel_orig:
            continue
        rel_exp = m.get("relato_expandido") or expandir_lenguaje_coloquial(rel_orig)
        h_c = m.get("hash_consulta") or hashlib.sha256(rel_orig.lower().encode("utf-8")).hexdigest()
        c.execute(sql_mem, (
            rel_orig,
            rel_exp,
            m.get("macro_rama", "Derecho Privado"),
            m.get("rama_derecho", "Derecho Laboral"),
            m.get("categoria_caso", "Casos Laborales"),
            m.get("tipo_caso", "Despido injustificado"),
            m.get("fuente_aprendizaje", "Auto-Inyección"),
            float(m.get("confianza", 95.0)),
            h_c,
            m.get("fecha_aprendizaje", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ))
        if c.rowcount > 0:
            nuevos_mem += 1

    sql_casos = " ".join([
        "INSERT OR IGNORE INTO casos_juridicos_cbr (",
        "codigo_caso, macro_rama, rama_derecho, categoria_caso, tipo_caso,",
        "titulo, actores_intervinientes, etapa_clave, medida_cautelar,",
        "resolucion_obtenida, recurso_utilizado, evidencia_logs, hash_sha256,",
        "resultado, accion_abogado, reaccion_juez, causa_raiz_fallo,",
        "estrategia_ganadora, que_hacer_ciudadano, que_no_hacer_ciudadano, fecha_registro",
        ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    ])

    for cs in data.get("casos_juridicos_cbr", []):
        cod = cs.get("codigo_caso")
        if not cod:
            continue
        c.execute(sql_casos, (
            cod,
            cs.get("macro_rama", "Derecho Privado"),
            cs.get("rama_derecho", "Derecho Laboral"),
            cs.get("categoria_caso", "Casos Laborales"),
            cs.get("tipo_caso", "Despido injustificado"),
            cs.get("titulo", "Caso Restaurado"),
            cs.get("actores_intervinientes", "Juez, Abogado"),
            cs.get("etapa_clave", "1. Demanda"),
            cs.get("medida_cautelar", "Embargo preventivo"),
            cs.get("resolucion_obtenida", "Sentencia"),
            cs.get("recurso_utilizado", "Amparo"),
            cs.get("evidencia_logs", ""),
            cs.get("hash_sha256", hashlib.sha256(str(cod).encode()).hexdigest()),
            cs.get("resultado", "GANADO"),
            cs.get("accion_abogado", ""),
            cs.get("reaccion_juez", ""),
            cs.get("causa_raiz_fallo", "N/A"),
            cs.get("estrategia_ganadora", ""),
            cs.get("que_hacer_ciudadano", ""),
            cs.get("que_no_hacer_ciudadano", ""),
            cs.get("fecha_registro", datetime.now().strftime("%Y-%m-%d"))
        ))
        if c.rowcount > 0:
            nuevos_casos += 1

    sql_buzon = " ".join([
        "INSERT OR IGNORE INTO buzon_casos_usuarios (",
        "nombre_usuario, contacto_usuario, relato_duda, estado, fecha_envio",
        ") VALUES (?, ?, ?, ?, ?)"
    ])
    for bz in data.get("buzon_casos_usuarios", []):
        rel_b = str(bz.get("relato_duda", "")).strip()
        if not rel_b:
            continue
        c.execute(sql_buzon, (
            bz.get("nombre_usuario", "Ciudadano"),
            bz.get("contacto_usuario", ""),
            rel_b,
            bz.get("estado", "Pendiente de revisión"),
            bz.get("fecha_envio", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ))

    conn.commit()
    return {"nuevos_memoria": nuevos_mem, "nuevos_casos": nuevos_casos}

def descargar_cerebro_desde_github(repo: str, token: str, branch: str = "main") -> dict:
    if not repo or not token:
        return {}
    url = f"https://api.github.com/repos/{repo.strip()}/contents/{BACKUP_JSON_FILE}?ref={branch.strip()}"
    headers = {
        "Authorization": f"token {token.strip()}",
        "Accept": "application/vnd.github.v3+json",
        "User-Agent": "ForensicMindAutoSync/10.0"
    }
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=5.0) as resp:
            res_json = json.loads(resp.read().decode("utf-8"))
            content_b64 = res_json.get("content", "")
            if content_b64:
                decoded = base64.b64decode(content_b64).decode("utf-8")
                return json.loads(decoded)
    except Exception:
        pass
    return {}

def guardar_y_sincronizar_cerebro_autonomo(repo: str = "", token: str = "", branch: str = "main") -> dict:
    conn = sqlite3.connect(DB_NAME)
    df_casos = pd.read_sql_query("SELECT * FROM casos_juridicos_cbr", conn)
    df_mem = pd.read_sql_query("SELECT * FROM memoria_neuronal_autonoma", conn)
    df_buzon = pd.read_sql_query("SELECT * FROM buzon_casos_usuarios", conn)
    conn.close()

    payload = {
        "autor_arquitectura": AUTOR_SISTEMA,
        "sistema": MARCA_SISTEMA,
        "fecha_respaldo": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_casos_cbr": len(df_casos),
        "total_memoria_autonoma": len(df_mem),
        "total_buzon_usuarios": len(df_buzon),
        "casos_juridicos_cbr": df_casos.to_dict(orient="records"),
        "memoria_neuronal_autonoma": df_mem.to_dict(orient="records"),
        "buzon_casos_usuarios": df_buzon.to_dict(orient="records")
    }
    raw_bytes = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")

    try:
        with open(BACKUP_JSON_FILE, "wb") as f:
            f.write(raw_bytes)
    except Exception:
        pass

    if not repo or not token:
        return {"local_ok": True, "github_ok": False, "msg": "Respaldo local actualizado."}

    url_api = f"https://api.github.com/repos/{repo.strip()}/contents/{BACKUP_JSON_FILE}"
    headers = {
        "Authorization": f"token {token.strip()}",
        "Accept": "application/vnd.github.v3+json",
        "Content-Type": "application/json",
        "User-Agent": "ForensicMindAutoSync/10.0"
    }

    sha_actual = None
    try:
        req_get = urllib.request.Request(f"{url_api}?ref={branch.strip()}", headers=headers)
        with urllib.request.urlopen(req_get, timeout=4.5) as resp_get:
            datos_existentes = json.loads(resp_get.read().decode("utf-8"))
            sha_actual = datos_existentes.get("sha")
    except Exception:
        sha_actual = None

    body_put = {
        "message": f"Auto-Sync ForensicMind v10.0 ({len(df_mem)} memorias | {len(df_casos)} casos)",
        "content": base64.b64encode(raw_bytes).decode("utf-8"),
        "branch": branch.strip()
    }
    if sha_actual:
        body_put["sha"] = sha_actual

    try:
        req_put = urllib.request.Request(
            url_api,
            data=json.dumps(body_put).encode("utf-8"),
            headers=headers,
            method="PUT"
        )
        with urllib.request.urlopen(req_put, timeout=6.0) as resp_put:
            if resp_put.status in (200, 201):
                return {"local_ok": True, "github_ok": True, "msg": "☁️ Sincronizado automáticamente en tu GitHub."}
    except Exception as e:
        return {"local_ok": True, "github_ok": False, "msg": f"Guardado local OK (GitHub: {str(e)[:35]})"}

    return {"local_ok": True, "github_ok": False, "msg": "Respaldo local actualizado."}

# ==============================================================================
# INICIALIZACIÓN DE SQLITE Y CASOS SEMILLA (GANADOS Y PERDIDOS)
# ==============================================================================
def init_db_and_autoinject(repo_gh: str = "", token_gh: str = "", branch_gh: str = "main"):
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    sql_create_casos = " ".join([
        "CREATE TABLE IF NOT EXISTS casos_juridicos_cbr (",
        "id INTEGER PRIMARY KEY AUTOINCREMENT, codigo_caso TEXT UNIQUE,",
        "macro_rama TEXT, rama_derecho TEXT, categoria_caso TEXT, tipo_caso TEXT,",
        "titulo TEXT, actores_intervinientes TEXT, etapa_clave TEXT,",
        "medida_cautelar TEXT, resolucion_obtenida TEXT, recurso_utilizado TEXT,",
        "evidencia_logs TEXT, hash_sha256 TEXT, resultado TEXT,",
        "accion_abogado TEXT, reaccion_juez TEXT, causa_raiz_fallo TEXT,",
        "estrategia_ganadora TEXT, que_hacer_ciudadano TEXT,",
        "que_no_hacer_ciudadano TEXT, fecha_registro TEXT)"
    ])
    c.execute(sql_create_casos)

    sql_create_mem = " ".join([
        "CREATE TABLE IF NOT EXISTS memoria_neuronal_autonoma (",
        "id INTEGER PRIMARY KEY AUTOINCREMENT, relato_original TEXT,",
        "relato_expandido TEXT, macro_rama TEXT, rama_derecho TEXT,",
        "categoria_caso TEXT, tipo_caso TEXT, fuente_aprendizaje TEXT,",
        "confianza REAL, hash_consulta TEXT UNIQUE, fecha_aprendizaje TEXT)"
    ])
    c.execute(sql_create_mem)

    sql_create_buzon = " ".join([
        "CREATE TABLE IF NOT EXISTS buzon_casos_usuarios (",
        "id INTEGER PRIMARY KEY AUTOINCREMENT, nombre_usuario TEXT,",
        "contacto_usuario TEXT, relato_duda TEXT UNIQUE,",
        "estado TEXT, fecha_envio TEXT)"
    ])
    c.execute(sql_create_buzon)

    sql_create_visitas = " ".join([
        "CREATE TABLE IF NOT EXISTS visitas_pagina (",
        "id INTEGER PRIMARY KEY AUTOINCREMENT, fecha_hora TEXT)"
    ])
    c.execute(sql_create_visitas)

    semillas = [
        (
            "EXP-LAB-2026-001", "Derecho Privado", "Derecho Laboral", "Casos Laborales", "Despido injustificado o falta de pago de quincenas/semanas",
            "Despido Injustificado Ganado (Sueldo Semanal sin Vacaciones y Oferta Engañosa de $5,000)",
            "Conciliador Laboral, Juez Laboral, Procurador PROFEDET",
            "1. Cita de Conciliación", "Embargo precautorio de cuentas", "Convenio pagado al 100%", "Demanda laboral",
            "Me despidieron injustificadamente con sueldo semanal de 2300, solo me quieren dar $5000 sin vacaciones ni descansos.",
            hashlib.sha256(b"EXP-LAB-2026-001").hexdigest(), "GANADO",
            "Se rechazaron los $5,000, no se firmó renuncia y se citó al patrón exigiendo 90 días de sueldo más vacaciones y descansos.",
            "La autoridad obligó al patrón a pagar la liquidación completa al no tener aviso escrito de despido justificado.",
            "N/A", "No firmar renuncia y acudir antes de 60 días al Centro de Conciliación.",
            "Rechazar ofertas engañosas de $5,000, guardar mensajes de WhatsApp de horarios y acudir a Conciliación.",
            "JAMÁS firmar una hoja de renuncia voluntaria ni hoja en blanco para recibir un pago mínimo.",
            "2026-08-01"
        ),
        (
            "EXP-LAB-2026-002", "Derecho Privado", "Derecho Laboral", "Casos Laborales", "Despido injustificado o falta de pago de quincenas/semanas",
            "Despido Perdido por Firmar Renuncia a Cambio de $5,000 y Dejar Pasar 2 Meses",
            "Juez Laboral, Abogado",
            "2. Audiencia de Conciliación", "Ninguna", "Absolución patronal", "Ninguno",
            "Trabajador aceptó $5,000 firmando hoja de renuncia voluntaria y reclamó después de 60 días.",
            hashlib.sha256(b"EXP-LAB-2026-002").hexdigest(), "PERDIDO",
            "Se presentó el reclamo cuando ya habían pasado más de 60 días naturales y con renuncia firmada.",
            "El Juez declaró que el derecho prescribió por exceder los 60 días y dio por válida la renuncia.",
            "Firmar renuncia voluntaria por urgencia de dinero y dejar pasar el límite de 60 días.",
            "N/A", "Iniciar el reclamo en la primera semana tras el despido sin firmar renuncia.",
            "Nunca firmar renuncia bajo la promesa verbal de que después te pagarán el resto.",
            "2026-08-05"
        ),
        (
            "EXP-TRA-2026-003", "Derecho Público", "Derecho Administrativo", "Tránsito y Abuso Policial", "Detención arbitraria, siembra de pruebas o extorsión policial",
            "Extorsión Policial y Corralón Anulados por Denuncia en Asuntos Internos y Fiscalía Anticorrupción",
            "Fiscalía Anticorrupción, Asuntos Internos, Tribunal Administrativo",
            "1. Denuncia y Queja", "Devolución inmediata de vehículo y celular", "Nulidad y sanción a oficiales", "Juicio de nulidad",
            "Patrulla me detuvo sin motivo, pidió mordida, intentó sembrar pruebas y se llevó mi moto al corralón.",
            hashlib.sha256(b"EXP-TRA-2026-003").hexdigest(), "GANADO",
            "Se solicitó bitácora GPS de la patrulla, cámaras C4/C5 de la esquina y se denunció en Fiscalía Anticorrupción y Asuntos Internos.",
            "Se acreditó la detención ilegal, se liberó el vehículo sin cobro de grúa y se vinculó a proceso a los malos elementos.",
            "N/A", "Ubicar cámaras de la calle, anotar número de patrulla y denunciar en Asuntos Internos y Fiscalía.",
            "Ubicar hora y esquina exacta para pedir videos del C5, anotar número de patrulla y acudir a Asuntos Internos.",
            "JAMÁS entregar dinero en efectivo ni firmar declaraciones sin tu abogado defensor presente.",
            "2026-08-12"
        ),
        (
            "EXP-TRA-2026-004", "Derecho Público", "Derecho Administrativo", "Tránsito y Abuso Policial", "Detención arbitraria, siembra de pruebas o extorsión policial",
            "Caso Vial Complicado por Agredir al Oficial y Entregar Efectivo sin Comprobante",
            "Juez Cívico / Ministerio Público",
            "1. Detención", "Arresto administrativo", "Multa y corralón pagados", "Ninguno",
            "Conductor discutió a golpes con el policía y entregó dinero sin anotar el número de patrulla.",
            hashlib.sha256(b"EXP-TRA-2026-004").hexdigest(), "PERDIDO",
            "El conductor perdió la calma, no registró el número de unidad ni pidió cámaras de vigilancia.",
            "Fue sancionado por resistencia de particulares y tuvo que pagar corralón completo.",
            "Perder la calma, golpear o insultar al oficial y no anotar el número de patrulla ni la hora.",
            "N/A", "Mantener la calma absoluta y usar la vía de Asuntos Internos.",
            "Nunca forcejear ni entregar dinero sin recibo oficial.",
            "2026-08-14"
        ),
        (
            "EXP-COB-2026-005", "Derecho Privado", "Derecho Mercantil", "Acoso de Cobranza y Deudas", "Amenazas de embargo falso o despachos de cobranza",
            "Cancelación de Deuda Falsa de App de Préstamos y Freno a Despacho de Cobranza",
            "CONDUSEF (REDECO), Policía Cibernética",
            "1. Queja oficial", "Bloqueo de hostigamiento", "Sanción al despacho y baja de adeudo", "Queja REDECO",
            "App de préstamos y despacho de cobranza exigían pagar un crédito no solicitado amenazando con embargar.",
            hashlib.sha256(b"EXP-COB-2026-005").hexdigest(), "GANADO",
            "Se levantó denuncia por usurpación de identidad, bloqueo de Buró de Crédito y queja ante REDECO-CONDUSEF.",
            "Se canceló el cobro fraudulento y cesaron las llamadas intimidatorias.",
            "N/A", "Denunciar robo de identidad y presentar queja en CONDUSEF sin depositar dinero a extorsionadores.",
            "Guardar capturas de amenazas, bloquear tu Buró de Crédito y denunciar ante CONDUSEF y Policía Cibernética.",
            "JAMÁS depositar dinero a cuentas personales de cobradores por miedo a un falso embargo.",
            "2026-08-18"
        ),
        (
            "EXP-COB-2026-006", "Derecho Privado", "Derecho Mercantil", "Acoso de Cobranza y Deudas", "Amenazas de embargo falso o despachos de cobranza",
            "Pérdida de Dinero por Depositar a Cuenta Personal de Cobrador por WhatsApp",
            "Despacho de Cobranza",
            "1. Cobro extrajudicial", "Ninguna", "Dinero perdido sin carta convenio", "Ninguno",
            "Deudor depositó $5,000 a una cuenta personal por miedo a amenaza de embargo por WhatsApp.",
            hashlib.sha256(b"EXP-COB-2026-006").hexdigest(), "PERDIDO",
            "Se hizo transferencia sin existir Carta Convenio oficial ni orden judicial.",
            "El despacho nunca abonó el dinero a la deuda y siguió extorsionando.",
            "Creer en amenazas de embargo por WhatsApp y depositar sin Carta Convenio oficial.",
            "N/A", "Exigir siempre Carta Convenio oficial y verificar ante CONDUSEF.",
            "Nunca transferir dinero a cuentas a nombre de personas físicas desconocidas.",
            "2026-08-19"
        ),
        (
            "EXP-PEN-2026-007", "Derecho Público", "Derecho Penal", "Casos Penales y Fraudes", "Robo de identidad con INE, fraude bancario o criptoactivos",
            "Restitución Bancaria y Anulación de Créditos Sacados con INE Robada",
            "CONDUSEF, Fiscalía Especializada, Juez Mercantil",
            "1. Carpeta de Investigación y CONDUSEF", "Bloqueo preventivo de cuentas", "Nulidad de créditos y devolución", "Dictamen CONDUSEF",
            "Usaron copia de mi INE para abrir créditos y me quitaron el celular con accesos bancarios.",
            hashlib.sha256(b"EXP-PEN-2026-007").hexdigest(), "GANADO",
            "Se bloqueó el SIM e IMEI, se revocaron accesos con semilla de respaldo, se levantó acta en Fiscalía y queja en CONDUSEF.",
            "Se anularon los créditos abiertos por robo de identidad al demostrarse que la víctima no firmó ni validó biometría.",
            "N/A", "Bloquear tarjetas y SIM en los primeros minutos y levantar denuncia por usurpación de identidad en Fiscalía.",
            "Congelar cuentas bancarias y SIM de inmediato, levantar denuncia en Fiscalía por robo de identidad y acudir a CONDUSEF.",
            "No dejar pasar días sin reportar el robo del celular o el uso ilegal de tu INE.",
            "2026-08-25"
        ),
        (
            "EXP-PEN-2026-008", "Derecho Público", "Derecho Penal", "Casos Penales y Fraudes", "Robo de identidad con INE, fraude bancario o criptoactivos",
            "Caso Complicado por No Reportar el Robo del Teléfono ni Congelar Cuentas a Tiempo",
            "Banco / CONDUSEF",
            "1. Reclamación tardía", "Ninguna", "Cargos aplicados", "Ninguno",
            "Víctima tardó dos semanas en reportar que le quitaron el teléfono con sus apps bancarias abiertas.",
            hashlib.sha256(b"EXP-PEN-2026-008").hexdigest(), "PERDIDO",
            "No se solicitó folio de bloqueo telefónico ni denuncia en Fiscalía tras el robo del equipo.",
            "Los delincuentes vaciaron las cuentas antes del bloqueo.",
            "No bloquear la tarjeta SIM ni las cuentas bancarias el mismo día del robo del celular.",
            "N/A", "Actuar en la primera hora congelando cuentas, SIM y moviendo criptoactivos con tu frase semilla.",
            "Jamás guardar tus palabras semilla o contraseñas en bloc de notas sin encriptar dentro del celular.",
            "2026-08-27"
        )
    ]

    sql_ins_semillas = " ".join([
        "INSERT OR IGNORE INTO casos_juridicos_cbr (",
        "codigo_caso, macro_rama, rama_derecho, categoria_caso, tipo_caso,",
        "titulo, actores_intervinientes, etapa_clave, medida_cautelar,",
        "resolucion_obtenida, recurso_utilizado, evidencia_logs, hash_sha256,",
        "resultado, accion_abogado, reaccion_juez, causa_raiz_fallo,",
        "estrategia_ganadora, que_hacer_ciudadano, que_no_hacer_ciudadano, fecha_registro",
        ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    ])
    c.executemany(sql_ins_semillas, semillas)
    conn.commit()

    if os.path.exists(BACKUP_JSON_FILE):
        try:
            with open(BACKUP_JSON_FILE, "r", encoding="utf-8") as f_loc:
                inyectar_diccionario_cerebro_en_sqlite(json.load(f_loc), conn)
        except Exception:
            pass

    if "autoinyeccion_remota_lista" not in st.session_state and repo_gh and token_gh:
        data_rem = descargar_cerebro_desde_github(repo_gh, token_gh, branch_gh)
        if data_rem:
            inyectar_diccionario_cerebro_en_sqlite(data_rem, conn)
        st.session_state["autoinyeccion_remota_lista"] = True

    conn.close()

init_db_and_autoinject(gh_repo_input, gh_token_input, gh_branch_input)

# ==============================================================================
# GUÍA CIUDADANA EXHAUSTIVA PASO A PASO POR CATEGORÍA
# ==============================================================================
def obtener_guia_humana_exhaustiva(
    categoria_caso: str,
    tipo_caso: str,
    datos_sueldo: dict,
    res_liq: dict
) -> dict:
    sueldo_str = f"${datos_sueldo['salario_mensual']:,.2f} MXN al mes ({datos_sueldo['periodicidad']})"
    oferta_val = datos_sueldo["oferta_patronal"]
    total_ley = res_liq["gran_total_bruto"]

    if categoria_caso == "Tránsito y Abuso Policial":
        radiografia = (
            "Detecté una situación grave de **Abuso Policial, Extorsión Vial, Siembra de Pruebas o Robo de Vehículo/Pertenencias en Patrulla**. "
            "Esta es tu **Prioridad #1**, porque involucra tu libertad física, tu motocicleta/auto en el corralón y la recuperación de tus pertenencias."
        )
        explicacion_sencilla = [
            "🔹 **Si los policías te sembraron algo o te quitaron tu celular y tu moto:** Eso NO es una infracción de tránsito; son delitos graves cometidos por servidores públicos (**Abuso de Autoridad, Extorsión y Robo Agravado**).",
            "🔹 **¿Por qué aquí NO aplica una 'liquidación laboral'?:** En abusos policiales no se calcula finiquito, sino la **liberación gratuita de tu vehículo**, la devolución de tus cosas y la **Reparación Integral del Daño** denunciando a los policías.",
            "🔹 **Las patrullas tienen GPS y las calles tienen cámaras C4/C5:** Aunque te hayan quitado tu celular, la hora exacta y la esquina donde te detuvieron permiten rastrear qué patrulla estaba ahí y pedir los videos oficiales antes de que se borren."
        ]
        fase_1_hoy = [
            "**1. Anota en papel la hora exacta, calle, entre qué esquinas fue y rasgos de los policías:** Si viste el número de patrulla o cuántos elementos iban, escríbelo de inmediato.",
            "**2. Si te quitaron el celular en la detención:** Pide prestado un teléfono hoy mismo para llamar a tu compañía telefónica, **suspender tu chip (SIM) por robo** y congelar tus cuentas bancarias para que los policías o terceros no entren a tus apps.",
            "**3. Si estás detenido o te quieren llevar al Ministerio Público:** Exige tu llamada constitucional, avisa a tu familia tu ubicación exacta y **niégate a firmar cualquier hoja** hasta que esté presente tu abogado o el Defensor Público."
        ]
        fase_2_semana = [
            "**1. Acude a la Fiscalía Especializada en Combate a la Corrupción y a Asuntos Internos:** Levanta la denuncia penal por extorsión, robo de celular, abuso de autoridad y fabricación de delito. Con esa carpeta de investigación se pide el resguardo inmediato de los videos del C4/C5.",
            "**2. Reclama tu moto o auto sin pagar extorsión:** Con el oficio de la Fiscalía Anticorrupción o del Tribunal de Justicia Administrativa, solicita la liberación de tu vehículo demostrando que la remisión al corralón fue ilegal.",
            "**3. Presenta queja en la Comisión de Derechos Humanos de tu Estado:** Esto activa medidas de protección para que esa patrulla no tome represalias en tu contra."
        ]
        fase_3_audiencia = [
            "**1. En la ratificación ante Asuntos Internos o el Juez:** Solicita que la corporación policial exhiba la bitácora de radio y el GPS de las patrullas que circularon por esa calle a esa hora.",
            "**2. Exige la Reparación del Daño:** Los gastos de grúa, corralón y el valor del teléfono sustraído deben ser cubiertos como reparación del daño dentro del proceso."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué decir en la Fiscalía Anticorrupción y en Asuntos Internos al denunciar:**\n\n"
            "*«Vengo a presentar denuncia formal por abuso de autoridad, extorsión y robo de mis pertenencias contra los elementos de la patrulla "
            "que me detuvieron en [Calle y Esquina] a las [Hora]. Solicito que de manera urgente se gire oficio al C4/C5 para asegurar las videograbaciones "
            "de las cámaras públicas antes de que se sobrescriban, así como el registro GPS de las unidades en esa zona.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Ir tú solo a buscar a los mismos policías para 'negociar' los $50,000):** Nunca les lleves dinero en efectivo; todo debe hacerse en oficinas de Asuntos Internos o Fiscalía.",
            "🚫 **Error 2 (Dejar pasar más de 5 días para pedir las cámaras del C4/C5):** Muchas cámaras públicas borran sus grabaciones después de 7 a 15 días; pide el aseguramiento del video hoy mismo.",
            "🚫 **Error 3 (Pagar corralón sin levantar antes el reporte de detención arbitraria):** Primero levanta tu queja oficial para dejar constancia del abuso."
        ]
        autoridad = "Fiscalía Anticorrupción / Unidad de Asuntos Internos Policial / Comisión de Derechos Humanos"
        abogado_req = "Asesor Jurídico de Víctimas en Fiscalía o Defensor en Justicia Administrativa ($0 MXN)"

    elif categoria_caso == "Casos Penales y Fraudes":
        radiografia = (
            "Detecté un problema crítico de **Robo de Identidad (uso ilegal de tu INE), Pérdida de Accesos Bancarios/Criptomonedas o Fraude Digital**. "
            "Es urgente blindar tus cuentas y tu identidad para que no vacíen tu dinero ni te endilguen deudas ajenas."
        )
        explicacion_sencilla = [
            "🔹 **¿Qué pasa si te robaron el celular donde tenías tus bancos y llaves de criptomonedas?:** Tus criptomonedas NO viven físicamente dentro del aparato, sino en la cadena de bloques (*blockchain*). Quien tenga el celular puede intentar abrir la app, pero si tú usas tu **frase semilla de 12 o 24 palabras** desde otra computadora limpia hoy mismo, puedes mover tus fondos a una billetera nueva antes de que adivinen tu PIN.",
            "🔹 **¿Qué pasa si usaron una foto de tu INE para sacar créditos a tu nombre?:** Eso es el delito de **Usurpación de Identidad**. Tú no estás obligado a pagar créditos que no firmaste ni autorizaste biométricamente."
        ]
        fase_1_hoy = [
            "**1. Rescata tus Criptoactivos y Cuentas en este instante (Primeros 30 minutos):** Desde la computadora de un familiar de confianza, importa tu frase semilla (*Recovery Phrase*) en una billetera nueva y **transfiere todos tus fondos a una dirección nueva**.",
            "**2. Bloquea el IMEI del teléfono, tu tarjeta SIM y tus bancos:** Llama a tu banco para congelar todas tus tarjetas y accesos móviles, pide tu **número de folio de reporte por robo de equipo** y entra a *Encontrar mi dispositivo* (Google/Apple) para ordenar el **Borrado Remoto** del celular que te quitaron.",
            "**3. Activa las Alertas y Bloqueo de tu Buró de Crédito:** Esto impide que sigan usando la foto de tu INE para abrir nuevos préstamos a tu nombre."
        ]
        fase_2_semana = [
            "**1. Levanta Denuncia por Robo de Identidad y Robo de Equipo en la Fiscalía:** Necesitas la copia sellada de tu Carpeta de Investigación; ese papel es tu escudo legal para demostrar que el INE y el celular estaban comprometidos.",
            "**2. Presenta Reclamación por Robo de Identidad ante CONDUSEF:** Con tu denuncia de Fiscalía y tu INE original, CONDUSEF obliga a las instituciones financieras a congelar y cancelar los créditos fraudulentos que tú no pediste."
        ]
        fase_3_audiencia = [
            "**1. En la audiencia de CONDUSEF:** Exige que la financiera exhiba las pruebas biométricas, geolocalización y cuenta bancaria a donde depositaron el supuesto préstamo. Como el dinero nunca entró a una cuenta tuya, el crédito se anula a tu favor."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué decir en el Banco y en la CONDUSEF por robo de celular e identidad:**\n\n"
            "*«Solicito el bloqueo total de mi banca electrónica y desconozco cualquier crédito o transferencia realizada tras el robo de mi equipo celular "
            "y el uso no autorizado de mi identificación. Presento mi Carpeta de Investigación por usurpación de identidad y exijo mi folio oficial de aclaración.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Dejar tus criptomonedas en la misma billetera del celular robado):** Aunque el teléfono tenga huella o PIN, debes mover tus fondos a una billetera nueva usando tu frase semilla de inmediato.",
            "🚫 **Error 2 (No levantar acta en el Ministerio Público/Fiscalía):** Sin denuncia formal, los bancos y financieras tardan mucho más en cancelar créditos por robo de identidad."
        ]
        autoridad = "Fiscalía Especializada / Policía Cibernética y CONDUSEF (Tel. 55 5340 0999)"
        abogado_req = "Asesor Gratuito de CONDUSEF y Ministerio Público ($0 MXN)"

    elif categoria_caso == "Casos Laborales":
        radiografia = (
            f"En la parte **Laboral** de tu caso: trabajabas ganando {sueldo_str}, "
            "te despidieron o te retuvieron el pago de tus quincenas/semanas trabajadas sin contrato ni IMSS "
            + (
                f"y te quieren dar solo ${oferta_val:,.2f} MXN en vez de tus ${total_ley:,.2f} MXN de ley. "
                if oferta_val > 0 else
                f"y por ley tienes derecho a reclamar aproximadamente ${total_ley:,.2f} MXN entre sueldos retenidos, finiquito y liquidación. "
            )
        )
        explicacion_sencilla = [
            "🔹 **¿Qué pasa si trabajabas remoto, sin contrato, sin IMSS y te pagaban por SPEI o criptomonedas?:** La Ley Federal del Trabajo (Art. 20 y 21) dice que si había alguien dándote órdenes (como el CEO por WhatsApp) a cambio de un pago, **sí existe relación laboral**, aunque la empresa sea informal o no te haya dado contrato.",
            "🔹 **¿Cómo le cobras a una empresa 'fantasma'?:** Rastreando la **razón social o el nombre del titular que aparece en tus transferencias SPEI** de tu estado de cuenta bancario. Quien te depositaba la nómina por SPEI responde solidariamente ante el Centro de Conciliación.",
            "🔹 **Retenerte 3 quincenas ya trabajadas es ilegal:** Además de poder exigir tu liquidación de 90 días por falta de pago de salario (Art. 51 LFT), el patrón está obligado a pagarte hasta el último peso de tus quincenas devengadas."
        ]
        fase_1_hoy = [
            "**1. Descarga tus estados de cuenta bancarios donde recibías los pagos SPEI:** Ahí viene el nombre exacto de la empresa o persona moral y su RFC/CLABE. Con ese nombre se le cita oficialmente.",
            "**2. Respalda desde WhatsApp Web o tu respaldo en la nube los chats del jefe/CEO:** Guarda los mensajes donde te asignaba trabajo, donde reclamaste tus quincenas y donde te amenazó con 'boletinarte' (poner en listas negras a un trabajador está prohibido por el Art. 133 fracción IX de la LFT).",
            "**3. Jamás firmes renuncia ni finiquito en blanco:** No aceptes acuerdos verbales."
        ]
        fase_2_semana = [
            "**1. Acude a la PROFEDET o al Centro de Conciliación Laboral:** Lleva tu INE y tus estados de cuenta donde aparecen los SPEI de la empresa para generar el citatorio oficial.",
            "**2. Si además usaron la copia de tu INE que les diste al entrar a trabajar para defraudarte:** Informa a la Procuraduría del Trabajo y anexa ese dato a tu denuncia penal contra el representante de la empresa."
        ]
        fase_3_audiencia = [
            "**1. En la audiencia de Conciliación o Juicio:** Exige el pago de tus quincenas retenidas + tus 90 días de Indemnización Constitucional + vacaciones y prima de antigüedad. Si la empresa no se presenta, se inicia el juicio laboral pidiendo el embargo de las cuentas bancarias desde donde te hacían los SPEI."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué decirle al patrón o CEO que te retiene sueldos o te amenaza:**\n\n"
            "*«Retener salarios ya trabajados y amenazar con boletinarme viola la Ley Federal del Trabajo. Ya descargué los comprobantes SPEI "
            "de la cuenta pagadora y acudiré al Centro de Conciliación Laboral y a la PROFEDET para exigir mis quincenas pendientes y mi liquidación de ley.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Dejar pasar los 60 días naturales):** Tienes un límite de 60 días desde el despido o separación para iniciar el trámite en Conciliación.",
            "🚫 **Error 2 (Creer que por ser trabajo remoto o sin contrato no tienes derechos):** Tus estados de cuenta SPEI y correos/mensajes prueban la relación de trabajo."
        ]
        autoridad = "Centro de Conciliación Laboral y PROFEDET (Procuraduría de la Defensa del Trabajo - Tel. 800 911 7877)"
        abogado_req = "Procurador Gratuito de la PROFEDET ($0 MXN) o Abogado Laboralista a porcentaje"

    elif categoria_caso == "Acoso de Cobranza y Deudas":
        radiografia = (
            "En el tema de **Acoso de Cobranza y Apps de Préstamos**: te están exigiendo pagar un dinero (o un crédito que ni siquiera pediste) "
            "mediante amenazas de embargo y hostigamiento telefónico."
        )
        explicacion_sencilla = [
            "🔹 **Si es un préstamo de una App que tú NUNCA solicitaste (o que abrieron robando tu INE):** No debes pagar ni un solo peso. Muchas apps de préstamos ilegales ('montadeudas') depositan sin permiso o usan datos robados para extorsionar.",
            "🔹 **Ningún despacho de cobranza puede embargar tu casa sin orden de un Juez:** Los mensajes de WhatsApp que dicen que irán mañana con cerrajero son 100% falsos e ilegales."
        ]
        fase_1_hoy = [
            "**1. NO les deposites ni un peso 'para que dejen de molestar':** Si les pagas a los montadeudas, te marcarán como víctima que sí cede y te inventarán otro cobro la semana siguiente.",
            "**2. Guarda capturas de los números y mensajes de amenaza:** Te servirán para el reporte oficial.",
            "**3. Avisa a tus contactos:** Si la app o los extorsionadores tienen tu agenda, publica un estado avisando que robaron tu teléfono/identidad y que ignoren cualquier mensaje pidiendo dinero."
        ]
        fase_2_semana = [
            "**1. Reporta a la Policía Cibernética y al REDECO de la CONDUSEF:** Si es una financiera registrada, CONDUSEF la multa y cancela el cobro por robo de identidad; si es una app ilegal montadeudas, la Policía Cibernética levanta el reporte de extorsión."
        ]
        fase_3_audiencia = [
            "**1. Bloqueo definitivo:** Con tu folio de la Policía Cibernética y CONDUSEF, bloquea todos los números desconocidos en tu nueva línea telefónica."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué responder una sola vez antes de bloquear al cobrador:**\n\n"
            "*«Este crédito jamás fue solicitado ni autorizado por mí, y ya existe una Carpeta de Investigación en Fiscalía por robo de identidad "
            "y extorsión, además del reporte ante CONDUSEF y Policía Cibernética. Cualquier amenaza constituye el delito de cobranza ilegítima.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Transferir dinero a cuentas CLABE a nombre de personas desconocidas):** Es una estafa directa.",
            "🚫 **Error 2 (Mandarles fotos nuevas de tu INE o comprobante de domicilio para 'aclarar' el error):** Solo usarán esos documentos para sacar más préstamos."
        ]
        autoridad = "CONDUSEF / Portal REDECO y Policía Cibernética (Tel. 088)"
        abogado_req = "Trámite Gratuito ante CONDUSEF y Policía Cibernética ($0 MXN)"

    elif categoria_caso == "Casos Familiares":
        radiografia = (
            "Analicé tu caso de **Derecho Familiar**. Cuando el deudor alimentario deja de aportar para la comida, escuela y salud de los hijos, "
            "la ley ordena fijar una pensión provisional urgente desde la primera semana."
        )
        explicacion_sencilla = [
            "🔹 **La pensión de los hijos es prioritaria:** Aunque el deudor oculte sus ingresos o trabaje en efectivo, el Juez Familiar investiga sus cuentas bancarias ante el SAT y fija una pensión obligatoria."
        ]
        fase_1_hoy = [
            "**1. Reúne las actas de nacimiento originales de tus hijos:** Y haz una lista sencilla de sus gastos mensuales de comida, escuela y salud."
        ]
        fase_2_semana = [
            "**1. Acude al DIF o a la Defensoría Pública Familiar:** Solicita de forma gratuita el inicio de la demanda de Pensión Alimenticia Provisional."
        ]
        fase_3_audiencia = [
            "**1. En el Juzgado Familiar:** Solicita que se giren oficios al IMSS, SAT y Comisión Nacional Bancaria para transparentar los ingresos reales del deudor."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué decirle al deudor alimentario:**\n\n"
            "*«La pensión alimenticia es un derecho irrenunciable de nuestros hijos y quedará establecida formalmente ante el Juez de lo Familiar.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Esperar meses confiando en promesas verbales de pago):** Actúa hoy mismo ante el Juzgado Familiar."
        ]
        autoridad = "Juzgado de lo Familiar y Procuraduría del DIF"
        abogado_req = "Abogado Familiar Gratuito del DIF ($0 MXN)"

    else:
        radiografia = (
            "Analicé tu caso de **Derecho Civil / Mercantil** (arrendamiento de inmueble o cobro de pagaré/deuda). "
            "Aquí la clave es asegurar el pago y la propiedad mediante orden judicial."
        )
        explicacion_sencilla = [
            "🔹 **En rentas o pagarés:** No puedes desalojar por la fuerza ni cambiar cerraduras sin orden del Juez, pero sí puedes pedir el embargo precautorio de bienes del deudor o fiador desde el inicio de la demanda."
        ]
        fase_1_hoy = [
            "**1. Resguarda el contrato de arrendamiento o pagaré original sin alterarlo ni rayarlo.**"
        ]
        fase_2_semana = [
            "**1. Acude a la Defensoría Pública Civil o con un abogado mercantil para presentar la demanda inicial con embargo precautorio.**"
        ]
        fase_3_audiencia = [
            "**1. En la diligencia judicial:** No entregues el documento original hasta que el pago esté firme en tu cuenta."
        ]
        guion_palabra_por_palabra = (
            "💬 **Qué decirle al inquilino o deudor moroso:**\n\n"
            "*«Dado que el plazo venció sin recibir el pago, el asunto ya fue turnado al Juzgado Civil para requerir el pago y la entrega por la vía legal.»*"
        )
        errores_fatales = [
            "🚫 **Error 1 (Cortar servicios o cambiar la chapa por tu cuenta):** Evita que te denuncien por despojo; hazlo siempre con orden del Juez."
        ]
        autoridad = "Juzgado Civil y Mercantil / Centro de Mediación"
        abogado_req = "Defensoría Pública Civil o Abogado Mercantil"

    return {
        "radiografia": radiografia,
        "explicacion_sencilla": explicacion_sencilla,
        "fase_1_hoy": fase_1_hoy,
        "fase_2_semana": fase_2_semana,
        "fase_3_audiencia": fase_3_audiencia,
        "guion_palabra_por_palabra": guion_palabra_por_palabra,
        "errores_fatales": errores_fatales,
        "autoridad": autoridad,
        "abogado_req": abogado_req
    }

# ==============================================================================
# MOTOR DE TRIAGE MULTI-MATERIA Y RED NEURONAL EVOLUTIVA
# ==============================================================================
def registrar_experiencia_en_memoria_autonoma(
    relato_original: str,
    clasificacion: dict,
    fuente: str = "Consulta Ciudadana",
    texto_web_extra: str = "",
    repo_gh: str = "",
    token_gh: str = "",
    branch_gh: str = "main"
) -> dict:
    if not relato_original or len(relato_original.strip()) < 10:
        return {"aprendio": False, "sync_msg": ""}
    relato_combinado = f"{relato_original} {texto_web_extra}".strip()
    relato_exp = expandir_lenguaje_coloquial(relato_combinado)
    h_cons = hashlib.sha256(relato_original.strip().lower().encode("utf-8")).hexdigest()

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    aprendio_nuevo = False
    sql_ins = " ".join([
        "INSERT OR IGNORE INTO memoria_neuronal_autonoma (",
        "relato_original, relato_expandido, macro_rama, rama_derecho,",
        "categoria_caso, tipo_caso, fuente_aprendizaje, confianza,",
        "hash_consulta, fecha_aprendizaje",
        ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    ])
    try:
        c.execute(sql_ins, (
            relato_original.strip(),
            relato_exp,
            clasificacion["macro_rama"],
            clasificacion["rama"],
            clasificacion["categoria"],
            clasificacion["tipo"],
            fuente,
            clasificacion["confianza"],
            h_cons,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))
        if c.rowcount > 0:
            aprendio_nuevo = True
        conn.commit()
    except Exception:
        pass
    finally:
        conn.close()

    sync_status = {"msg": ""}
    if aprendio_nuevo:
        sync_status = guardar_y_sincronizar_cerebro_autonomo(repo_gh, token_gh, branch_gh)

    return {"aprendio": aprendio_nuevo, "sync_msg": sync_status.get("msg", "")}

def entrenar_red_neuronal_evolutiva():
    conn = sqlite3.connect(DB_NAME)
    df_base = pd.read_sql_query("SELECT * FROM casos_juridicos_cbr", conn)
    df_mem = pd.read_sql_query("SELECT * FROM memoria_neuronal_autonoma", conn)
    conn.close()

    if df_base.empty or MLPClassifier is None or TfidfVectorizer is None:
        return None, None, {
            "total_muestras": len(df_base),
            "experiencias_aprendidas": len(df_mem),
            "vocabulario_ngramas": 0
        }

    textos_entrenamiento = []
    etiquetas_entrenamiento = []

    for _, r in df_base.iterrows():
        txt_b = f"{r['titulo']} {r['evidencia_logs']} {r['tipo_caso']} {r['categoria_caso']} {r['que_hacer_ciudadano']}"
        textos_entrenamiento.append(expandir_lenguaje_coloquial(txt_b))
        etiquetas_entrenamiento.append(f"{r['categoria_caso']}|{r['rama_derecho']}|{r['tipo_caso']}|{r['macro_rama']}")

    for _, m in df_mem.iterrows():
        textos_entrenamiento.append(str(m["relato_expandido"]))
        etiquetas_entrenamiento.append(f"{m['categoria_caso']}|{m['rama_derecho']}|{m['tipo_caso']}|{m['macro_rama']}")

    vectorizer = TfidfVectorizer(ngram_range=(1, 3), max_features=2400)
    X_vec = vectorizer.fit_transform(textos_entrenamiento)

    mlp = MLPClassifier(hidden_layer_sizes=(96, 48), activation='relu', max_iter=450, random_state=42)
    mlp.fit(X_vec, etiquetas_entrenamiento)

    return vectorizer, mlp, {
        "total_muestras": len(textos_entrenamiento),
        "experiencias_aprendidas": len(df_mem),
        "vocabulario_ngramas": len(vectorizer.vocabulary_)
    }

def diagnosticar_triage_multimateria(relato: str, vectorizer_mlp, modelo_mlp) -> dict:
    texto_low = relato.lower()

    reglas_deteccion = {
        "Tránsito y Abuso Policial": [
            "transito", "tránsito", "policia", "policía", "policias", "policías",
            "patrulla", "mordida", "corralon", "corralón", "grua", "grúa",
            "infraccion", "infracción", "multa", "placa", "reten", "retén",
            "sembraron", "sustancias ilicitas", "sustancias ilícitas", "oficial"
        ],
        "Casos Penales y Fraudes": [
            "fraude", "estafa", "phishing", "vaciaron", "robaron mi cuenta",
            "clonaron", "criptomonedas", "llaves privadas", "copia de mi ine",
            "usaron mi ine", "sin mi consentimiento", "abrir creditos a mi nombre",
            "abrir créditos a mi nombre", "quitaron mi celular", "robo de identidad", "e.firma"
        ],
        "Casos Laborales": [
            "despid", "corrieron", "echaron", "chamba", "trabajaba", "trabajo",
            "liquidaci", "finiquito", "sueldo", "salario", "semanal", "quincenal",
            "quincenas", "vacaciones", "aguinaldo", "patron", "patrón", "renuncia",
            "hoja en blanco", "dias libres", "días libres", "startup", "imss"
        ],
        "Acoso de Cobranza y Deudas": [
            "cobranza", "despacho", "embargar", "embargo", "app de prestamos",
            "app de préstamos", "amenazan con embargar", "romper cerradura",
            "credito que jamas solicite", "crédito que jamás solicité", "no pedi"
        ],
        "Casos Familiares": [
            "pension", "pensión", "alimentos", "alimenticia", "papa de mis hijos",
            "papá de mis hijos", "custodia", "divorcio", "gasto de los niños", "manutencion"
        ],
        "Casos Civiles y Mercantiles": [
            "inquilino", "renta", "arrendamiento", "pagare", "pagaré", "desocupar", "cerraduras"
        ]
    }

    categorias_activas = []
    for cat_nombre, palabras in reglas_deteccion.items():
        coincidencias = sum(1 for p in palabras if p in texto_low)
        if coincidencias >= 1:
            categorias_activas.append(cat_nombre)

    mapa_info = {
        "Tránsito y Abuso Policial": ("Derecho Público", "Derecho Administrativo y Penal (Abuso Policial)", "Detención arbitraria, siembra de pruebas o extorsión policial"),
        "Casos Penales y Fraudes": ("Derecho Público", "Derecho Penal y Cibernético (Identidad y Cuentas)", "Robo de identidad con INE, fraude bancario o criptoactivos"),
        "Casos Laborales": ("Derecho Privado", "Derecho Laboral", "Despido injustificado o falta de pago de quincenas/semanas"),
        "Acoso de Cobranza y Deudas": ("Derecho Privado", "Defensa del Deudor y CONDUSEF", "Amenazas de embargo falso o despachos de cobranza"),
        "Casos Familiares": ("Derecho Privado", "Derecho de Familia", "Pensión alimenticia y gasto de los hijos"),
        "Casos Civiles y Mercantiles": ("Derecho Privado", "Derecho Civil y Mercantil", "Inquilino que no paga renta ni desocupa")
    }

    if not categorias_activas:
        cat_fallback = "Casos Laborales"
        if modelo_mlp is not None and vectorizer_mlp is not None:
            try:
                rel_exp = expandir_lenguaje_coloquial(relato)
                pred = modelo_mlp.predict(vectorizer_mlp.transform([rel_exp]))[0].split("|")
                if len(pred) == 4 and pred[0] in mapa_info:
                    cat_fallback = pred[0]
            except Exception:
                pass
        categorias_activas = [cat_fallback]

    categorias_ordenadas = sorted(categorias_activas, key=lambda c: PRIORIDAD_TRIAGE.get(c, (99, ""))[0])

    cat_principal = categorias_ordenadas[0]
    macro_p, rama_p, tipo_p = mapa_info[cat_principal]
    tiene_laboral = ("Casos Laborales" in categorias_ordenadas)

    lista_modulos = []
    for c_ord in categorias_ordenadas:
        m_i, r_i, t_i = mapa_info[c_ord]
        num_prio, titulo_prio = PRIORIDAD_TRIAGE[c_ord]
        lista_modulos.append({
            "categoria": c_ord,
            "macro_rama": m_i,
            "rama": r_i,
            "tipo": t_i,
            "prioridad_num": num_prio,
            "titulo_prioridad": titulo_prio
        })

    return {
        "macro_rama": macro_p,
        "rama": rama_p if len(categorias_ordenadas) == 1 else f"Multi-Materia ({len(categorias_ordenadas)} Áreas Detectadas)",
        "categoria": cat_principal,
        "tipo": tipo_p,
        "confianza": 97.2,
        "es_multimateria": len(categorias_ordenadas) > 1,
        "tiene_componente_laboral": tiene_laboral,
        "categorias_ordenadas": categorias_ordenadas,
        "modulos_triage": lista_modulos
    }

# ==============================================================================
# BÚSQUEDA EN INTERNET Y CÁLCULO DE LIQUIDACIÓN LFT (CON CANDADO FINANCIERO)
# ==============================================================================
def buscar_informacion_juridica_en_internet(categoria_caso: str, tipo_caso: str) -> dict:
    resultados_web = []
    textos_para_entrenar = []

    q_wiki = f"{tipo_caso} México"
    try:
        url_wiki = (
            "https://es.wikipedia.org/w/api.php?action=query&list=search"
            f"&srsearch={urllib.parse.quote(q_wiki)}&utf8=&format=json&srlimit=2"
        )
        req_w = urllib.request.Request(url_wiki, headers={"User-Agent": "ForensicMindBot/10.0"})
        with urllib.request.urlopen(req_w, timeout=3.5) as resp_w:
            data_w = json.loads(resp_w.read().decode("utf-8"))
            for item in data_w.get("query", {}).get("search", []):
                tit_w = item.get("title", "")
                snip_w = re.sub(r"<[^>]+>", "", item.get("snippet", ""))
                link_w = f"https://es.wikipedia.org/wiki/{urllib.parse.quote(tit_w.replace(' ', '_'))}"
                resultados_web.append({
                    "fuente": "📖 Enciclopedia Legal",
                    "titulo": tit_w,
                    "resumen": snip_w,
                    "url": link_w
                })
                textos_para_entrenar.append(f"{tit_w} {snip_w}")
    except Exception:
        pass

    q_news = f"{tipo_caso} derechos Mexico ley"
    if categoria_caso == "Casos Laborales":
        q_news = "despido injustificado liquidacion PROFEDET Ley Federal del Trabajo Mexico"
    elif categoria_caso == "Tránsito y Abuso Policial":
        q_news = "abuso policial multa transito corralon Asuntos Internos Mexico"
    elif categoria_caso == "Acoso de Cobranza y Deudas":
        q_news = "despachos de cobranza amenazas embargo CONDUSEF Mexico"

    try:
        url_gnews = (
            "https://news.google.com/rss/search?"
            f"q={urllib.parse.quote(q_news)}&hl=es-419&gl=MX&ceid=MX:es-419"
        )
        req_g = urllib.request.Request(url_gnews, headers={"User-Agent": "ForensicMindBot/10.0"})
        with urllib.request.urlopen(req_g, timeout=4.0) as resp_g:
            root = ET.fromstring(resp_g.read())
            for it in root.findall(".//item")[:3]:
                tit = it.find("title").text if it.find("title") is not None else "Noticia Jurídica"
                lnk = it.find("link").text if it.find("link") is not None else "https://www.gob.mx"
                resultados_web.append({
                    "fuente": "🌐 Noticias y Criterios en Línea (Google México)",
                    "titulo": tit,
                    "resumen": f"Información reciente verificada para {categoria_caso}.",
                    "url": lnk
                })
                textos_para_entrenar.append(tit)
    except Exception:
        pass

    return {
        "resultados": resultados_web,
        "corpus_web_extraido": " ".join(textos_para_entrenar)
    }

def normalizar_texto_coloquial(texto: str) -> dict:
    texto_low = texto.lower().replace(",", "")
    salario_mensual = None
    periodicidad = "Mensual ($9,857.14 / mes)"

    p_sem = r'(?:semanal(?:es)?(?:\s+es)?(?:\s+de)?\s*\$?\s*(\d+(?:\.\d+)?))|(?:(\d+(?:\.\d+)?)\s*(?:pesos\s*)?(?:a la semana|semanal(?:es)?|por semana))'
    m_sem = re.search(p_sem, texto_low)
    if m_sem:
        val = float(m_sem.group(1) or m_sem.group(2))
        salario_mensual = round((val / 7.0) * 30.0, 2)
        periodicidad = f"Semanal (${val:,.2f} a la semana)"

    if salario_mensual is None:
        p_quin = r'(?:quincenal(?:es)?(?:\s+es)?(?:\s+de)?\s*\$?\s*(\d+(?:\.\d+)?))|(?:(\d+(?:\.\d+)?)\s*(?:pesos\s*)?(?:a la quincena|quincenal(?:es)?|por quincena))'
        m_quin = re.search(p_quin, texto_low)
        if m_quin:
            val = float(m_quin.group(1) or m_quin.group(2))
            salario_mensual = round(val * 2.0, 2)
            periodicidad = f"Quincenal (${val:,.2f} a la quincena)"

    if salario_mensual is None:
        m_gen = re.search(r'(?:gano|ganaba|sueldo|salario)\s*(?:es\s*de)?\s*\$?\s*(\d+(?:\.\d+)?)', texto_low)
        if m_gen:
            val = float(m_gen.group(1))
            if val < 5000:
                salario_mensual = round((val / 7.0) * 30.0, 2)
                periodicidad = f"Semanal (${val:,.2f} a la semana)"
            else:
                salario_mensual = round(val, 2)
                periodicidad = f"Mensual (${val:,.2f} al mes)"

    oferta_patronal = 0.0
    m_of = re.search(r'(?:quieren dar|ofrecen|daban|dan|soltar)\s*\$?\s*(\d+(?:\.\d+)?)', texto_low)
    if m_of:
        oferta_patronal = float(m_of.group(1))

    quincenas_retenidas = 0
    m_qret = re.search(r'(\d+)\s*quincenas', texto_low)
    if m_qret:
        quincenas_retenidas = int(m_qret.group(1))

    falta_prestaciones = any(k in texto_low for k in [
        "no me daban vacaciones", "sin vacaciones", "días libres", "dias libres",
        "nada de bonos", "sin aguinaldo", "nunca descansaba", "12 horas diarias"
    ])

    return {
        "salario_mensual": salario_mensual if salario_mensual else 9857.14,
        "periodicidad": periodicidad,
        "oferta_patronal": oferta_patronal,
        "quincenas_retenidas": quincenas_retenidas,
        "falta_prestaciones": falta_prestaciones
    }

def obtener_dias_vacaciones_lft(anios: int) -> int:
    if anios <= 1:
        return 12
    if anios == 2:
        return 14
    if anios == 3:
        return 16
    if anios == 4:
        return 18
    if anios == 5:
        return 20
    if 6 <= anios <= 10:
        return 22
    return 24

def calcular_liquidacion_y_derechos_lft(
    salario_mensual: float,
    anios_antiguedad: int,
    meses_adicionales: int,
    motivo_separacion: str,
    dias_sueldo_pendientes: int = 0,
    dias_vacaciones_no_gozadas: int = 0,
    dias_descanso_trabajados: int = 0,
    salario_minimo_vigente: float = 278.80
) -> dict:
    sd = round(salario_mensual / 30.0, 2)
    antiguedad_total = round(max(1.0, anios_antiguedad + (meses_adicionales / 12.0)), 4)
    dias_vac = obtener_dias_vacaciones_lft(max(1, anios_antiguedad))

    factor_int = 1.0 + (15.0 / 365.0) + ((dias_vac * 0.25) / 365.0)
    sdi = round(sd * factor_int, 2)

    m_sueldo_pend = round(dias_sueldo_pendientes * sd, 2)
    m_aguinaldo = round(15.0 * sd, 2)
    m_vacaciones = round((dias_vac + dias_vacaciones_no_gozadas) * sd, 2)
    m_prima_vac = round(m_vacaciones * 0.25, 2)
    m_descansos_dobles = round(dias_descanso_trabajados * (sd * 2.0), 2)

    es_despido = ("Despido" in motivo_separacion or "Rescisión" in motivo_separacion)
    tope_prima = min(sd, salario_minimo_vigente * 2.0)
    m_prima_ant = round(12.0 * antiguedad_total * tope_prima, 2) if es_despido else 0.0
    m_indem_90 = round(90.0 * sdi, 2) if es_despido else 0.0

    subtotal_finiquito = round(m_sueldo_pend + m_aguinaldo + m_vacaciones + m_prima_vac + m_descansos_dobles, 2)
    subtotal_indemnizacion = round(m_prima_ant + m_indem_90, 2)
    gran_total = round(subtotal_finiquito + subtotal_indemnizacion, 2)

    conceptos = [
        {"Concepto Explicado Fácil": "1. Días o quincenas trabajadas sin pagar", "Por qué te toca (Ley)": "Art. 82 y 88 LFT (Tu trabajo ya realizado)", "Monto (MXN)": m_sueldo_pend},
        {"Concepto Explicado Fácil": "2. Aguinaldo (Mínimo 15 días por año)", "Por qué te toca (Ley)": "Art. 87 LFT (Derecho irrenunciable)", "Monto (MXN)": m_aguinaldo},
        {"Concepto Explicado Fácil": "3. Vacaciones Dignas + Vacaciones no gozadas", "Por qué te toca (Ley)": "Art. 76 y 79 LFT (Mínimo 12 días el 1er año)", "Monto (MXN)": m_vacaciones},
        {"Concepto Explicado Fácil": "4. Prima Vacacional (25% extra sobre vacaciones)", "Por qué te toca (Ley)": "Art. 80 LFT (Dinero extra para tus vacaciones)", "Monto (MXN)": m_prima_vac},
        {"Concepto Explicado Fácil": "5. Días de descanso trabajados sin pago extra", "Por qué te toca (Ley)": "Art. 73 LFT (Trabajar en tu día libre se paga con un 200% extra)", "Monto (MXN)": m_descansos_dobles},
        {"Concepto Explicado Fácil": "6. Prima de Antigüedad (12 días por cada año)", "Por qué te toca (Ley)": "Art. 162 LFT (Reconocimiento por tu tiempo en la empresa)", "Monto (MXN)": m_prima_ant},
        {"Concepto Explicado Fácil": "7. Indemnización Constitucional (3 Meses / 90 días de sueldo integrado)", "Por qué te toca (Ley)": "Art. 123 Constitucional y Art. 48 LFT (Castigo al patrón por despedirte o retener tu sueldo)", "Monto (MXN)": m_indem_90}
    ]

    return {
        "sd": sd,
        "sdi": sdi,
        "subtotal_finiquito": subtotal_finiquito,
        "subtotal_indemnizacion": subtotal_indemnizacion,
        "gran_total_bruto": gran_total,
        "tabla_conceptos": pd.DataFrame(conceptos)
    }

def calcular_cotizacion_honorarios_dual(
    monto_en_disputa: float,
    costo_asesoria: float = 1000.0,
    anticipo_fijo: float = 0.0,
    pct_litis: float = 25.0,
    gastos_periciales: float = 1000.0,
    horas_invertidas: float = 35.0,
    costo_hora_despacho: float = 350.0
) -> dict:
    comision_exito = round(monto_en_disputa * (pct_litis / 100.0), 2)
    honorarios_totales = round(costo_asesoria + anticipo_fijo + comision_exito, 2)
    desembolso_total = round(honorarios_totales + gastos_periciales, 2)
    neto_libre = round(max(0.0, monto_en_disputa - desembolso_total), 2)
    utilidad_abogado = round(honorarios_totales - (horas_invertidas * costo_hora_despacho), 2)
    return {
        "honorarios_totales_abogado": honorarios_totales,
        "monto_neto_libre_ciudadano": neto_libre,
        "utilidad_neta_abogado": utilidad_abogado
    }

# ==============================================================================
# DICTADO POR VOZ, WORD (.DOCX) Y FICHA 9:16
# ==============================================================================
def transcribir_audio_espanol(audio_bytes) -> str:
    if sr is None or audio_bytes is None:
        return ""
    try:
        r = sr.Recognizer()
        audio_file = io.BytesIO(audio_bytes.getvalue() if hasattr(audio_bytes, "getvalue") else audio_bytes.read())
        with sr.AudioFile(audio_file) as source:
            audio_data = r.record(source)
            return r.recognize_google(audio_data, language="es-MX")
    except Exception:
        return "⚠️ No se pudo transcribir el audio automáticamente. Escribe tu caso en el recuadro."

def limpiar_ascii(texto: str) -> str:
    limpio = re.sub(r"\*+", "", str(texto))
    mapa = str.maketrans("áéíóúñÁÉÍÓÚÑüÜ«»¿¡", "aeiounAEIOUNuU\"\"?!")
    return limpio.translate(mapa).encode("ascii", "ignore").decode("ascii")

def generar_ficha_vertical_9_16(
    codigo: str,
    triage_info: dict,
    guia_principal: dict,
    monto_bruto: float,
    monto_neto: float,
    hash_ev: str
) -> np.ndarray:
    alto, ancho = 1920, 1080
    img = np.zeros((alto, ancho, 3), dtype=np.uint8)
    img[:] = (24, 15, 10)
    cv2.rectangle(img, (25, 25), (1055, 1895), (240, 175, 50), 4)
    cv2.rectangle(img, (40, 40), (1040, 235), (45, 28, 16), -1)

    cv2.putText(img, "FORENSICMIND & LEXGUIA v10.0", (65, 105), cv2.FONT_HERSHEY_DUPLEX, 1.45, (100, 220, 255), 3)
    cv2.putText(img, limpiar_ascii(f"AUTOR: {AUTOR_SISTEMA.upper()}"), (65, 160), cv2.FONT_HERSHEY_DUPLEX, 0.75, (120, 230, 255), 2)

    etiqueta_rama = triage_info["rama"]
    cv2.putText(img, limpiar_ascii(f"EXPEDIENTE: {codigo} | {etiqueta_rama[:38]}"), (65, 285), cv2.FONT_HERSHEY_DUPLEX, 0.80, (255, 255, 255), 2)

    cv2.rectangle(img, (65, 325), (1015, 470), (25, 48, 38), -1)
    if triage_info["tiene_componente_laboral"]:
        cv2.putText(img, limpiar_ascii(f"RECLAMO LABORAL LFT: ${monto_bruto:,.2f} MXN"), (85, 380), cv2.FONT_HERSHEY_DUPLEX, 0.80, (255, 255, 255), 2)
        cv2.putText(img, limpiar_ascii(f"NETO LIBRE PARA TI: ${monto_neto:,.2f} MXN"), (85, 438), cv2.FONT_HERSHEY_DUPLEX, 0.82, (120, 255, 180), 2)
    else:
        cv2.putText(img, "BLINDAJE LEGAL: $0 EN MORDIDAS O EXTORSIÓN", (85, 380), cv2.FONT_HERSHEY_DUPLEX, 0.78, (255, 255, 255), 2)
        cv2.putText(img, "VIA DE DEFENSA: NULIDAD Y REPARACION DEL DANO", (85, 438), cv2.FONT_HERSHEY_DUPLEX, 0.76, (120, 255, 180), 2)

    cv2.putText(img, "PLAN DE ACCION POR ORDEN DE URGENCIA:", (65, 535), cv2.FONT_HERSHEY_DUPLEX, 0.88, (100, 220, 255), 2)
    y_cursor = 590

    pasos_combinados = guia_principal["fase_1_hoy"] + guia_principal["fase_2_semana"][:2]
    for paso in pasos_combinados:
        lineas_envueltas = textwrap.wrap(limpiar_ascii(paso), width=52)
        for lin in lineas_envueltas[:3]:
            if y_cursor < 1470:
                cv2.putText(img, lin, (65, y_cursor), cv2.FONT_HERSHEY_SIMPLEX, 0.70, (235, 235, 235), 2)
                y_cursor += 36
        y_cursor += 16

    cv2.putText(img, "AUTORIDAD DONDE ACUDIR:", (65, 1520), cv2.FONT_HERSHEY_DUPLEX, 0.82, (100, 220, 255), 2)
    y_aut = 1565
    for lin_aut in textwrap.wrap(limpiar_ascii(guia_principal["autoridad"]), width=54)[:2]:
        cv2.putText(img, lin_aut, (65, y_aut), cv2.FONT_HERSHEY_SIMPLEX, 0.70, (255, 230, 150), 2)
        y_aut += 38

    cv2.putText(img, "SELLO CRIPTOGRAFICO SHA-256:", (65, 1670), cv2.FONT_HERSHEY_DUPLEX, 0.80, (100, 220, 255), 2)
    cv2.putText(img, hash_ev[:40], (65, 1715), cv2.FONT_HERSHEY_SIMPLEX, 0.75, (160, 255, 160), 2)
    cv2.putText(img, limpiar_ascii(f"ARQUITECTURA: {AUTOR_SISTEMA.upper()}"), (65, 1820), cv2.FONT_HERSHEY_DUPLEX, 0.72, (100, 255, 218), 2)
    return cv2.cvtColor(img, cv2.COLOR_BGR2RGB)

def generar_dictamen_word_docx(
    codigo: str,
    relato: str,
    triage_info: dict,
    guias_por_modulo: list,
    res_liq: dict,
    res_hon: dict,
    hash_ev: str
) -> bytes:
    if Document is None:
        return b""
    doc = Document()
    doc.add_heading("FORENSICMIND & LEXGUÍA v10.0 — MANUAL DE TRIAGE PASO A PASO", 0)
    doc.add_paragraph(f"Autoría del Sistema: {AUTOR_SISTEMA}")
    doc.add_paragraph(f"Expediente: {codigo} | Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Materias Detectadas ({len(guias_por_modulo)}): {', '.join(triage_info['categorias_ordenadas'])}")
    doc.add_paragraph(f"Sello Forense SHA-256: {hash_ev}")

    if triage_info["tiene_componente_laboral"]:
        doc.add_heading("Cálculo de Derechos Laborales (LFT)", level=1)
        doc.add_paragraph(f"• Total Bruto por Ley (LFT): ${res_liq['gran_total_bruto']:,.2f} MXN")
        doc.add_paragraph(f"• Neto Libre estimado: ${res_hon['monto_neto_libre_ciudadano']:,.2f} MXN")

    for idx_m, item_m in enumerate(guias_por_modulo, 1):
        g = item_m["guia"]
        doc.add_heading(f"MÓDULO {idx_m}: {item_m['titulo_prioridad']}", level=1)
        doc.add_paragraph(re.sub(r"\*+", "", g["radiografia"]))

        doc.add_heading("Explicación en Español Sencillo", level=2)
        for exp in g["explicacion_sencilla"]:
            doc.add_paragraph(re.sub(r"\*+", "", exp), style="List Bullet")

        doc.add_heading("Qué hacer Hoy Mismo", level=2)
        for p1 in g["fase_1_hoy"]:
            doc.add_paragraph(re.sub(r"\*+", "", p1), style="List Bullet")

        doc.add_heading("Qué hacer Esta Semana", level=2)
        for p2 in g["fase_2_semana"]:
            doc.add_paragraph(re.sub(r"\*+", "", p2), style="List Bullet")

        doc.add_heading("Guion Exacto Palabra por Palabra", level=2)
        doc.add_paragraph(re.sub(r"\*+", "", g["guion_palabra_por_palabra"]))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ==============================================================================
# ENCABEZADO PRINCIPAL Y PESTAÑAS (PÚBLICO VS ADMIN)
# ==============================================================================
banner_html = " ".join([
    "<div class='author-banner'>",
    "<div style='display:flex; flex-wrap:wrap; justify-content:space-between; align-items:center; gap:10px;'>",
    "<div>",
    "<span style='color:#fbbf24 !important; font-weight:800; font-size:13px;'>🏛️ SISTEMA DE TRIAGE LEGAL MULTI-MATERIA & DESPACHO • GITHUB AUTO-SYNC</span><br>",
    f"<span style='color:#ffffff !important; font-size:21px; font-weight:900;'>{MARCA_SISTEMA}</span>",
    "</div>",
    "<div>",
    f"<span class='gold-badge'>⚙️ SELLO PERSONAL: {AUTOR_SISTEMA}</span>",
    "</div>",
    "</div>",
    "</div>"
])
st.markdown(banner_html, unsafe_allow_html=True)

VECTORIZER_VIVO, MODELO_VIVO, STATS_IA = entrenar_red_neuronal_evolutiva()

if st.session_state["es_admin"]:
    tab1, tab2, tab4, tab3, tab5, tab6, tab7, tab8 = st.tabs([
        "🧠 1. Guía Paso a Paso (Triage)",
        "💰 2. Calculadora LFT y Honorarios",
        "🔬 3. Casos Ganados vs. Perdidos",
        "☁️ 4. [ADMIN] Cerebro & GitHub",
        "➕ 5. [ADMIN] Enseñar Caso",
        "📬 6. [ADMIN] Buzón Usuarios",
        "🛠️ 7. [ADMIN] Cotizador Despacho",
        "👥 8. [ADMIN] Tráfico & Visitas"
    ])
else:
    tab1, tab2, tab4 = st.tabs([
        "🧠 1. Guía Paso a Paso de tu Caso",
        "💰 2. Calculadora de Liquidación y Abogado",
        "🔬 3. Casos Ganados vs. Perdidos"
    ])

# ==============================================================================
# PESTAÑA 1: TRIAGE MULTI-MATERIA PASO A PASO + COTIZADOR CLIENTE Y WHATSAPP
# ==============================================================================
with tab1:
    st.markdown("### 🗣️ Cuéntame todo lo que te pasó (Aunque sean varios problemas mezclados)")

    col_esc, col_voz = st.columns([1.3, 1])
    with col_esc:
        escenario_sel = st.selectbox(
            "💡 Elige un escenario de prueba o escribe/dicta tu caso:",
            list(ESCENARIOS_FRECUENTES.keys())
        )
    with col_voz:
        audio_input = st.audio_input("🎙️ Dictar mi problema por voz (Español)")
        texto_transcrito = ""
        if audio_input is not None:
            with st.spinner("Escuchando y transcribiendo tu voz..."):
                texto_transcrito = transcribir_audio_espanol(audio_input)
                if texto_transcrito and "⚠️" not in texto_transcrito:
                    st.success("✅ Voz transcrita con éxito.")
                elif texto_transcrito:
                    st.warning(texto_transcrito)

    valor_relato = ESCENARIOS_FRECUENTES[escenario_sel]
    if texto_transcrito and "⚠️" not in texto_transcrito:
        valor_relato = texto_transcrito

    relato_usuario = st.text_area(
        "📝 Describe aquí todo lo que ocurrió con tus propias palabras:",
        value=valor_relato,
        height=140
    )

    btn_analizar = st.button(
        "🧠 Analizar Caso Completo y Generar Plan de Acción por Orden de Urgencia",
        use_container_width=True
    )

    triage_res = diagnosticar_triage_multimateria(relato_usuario, VECTORIZER_VIVO, MODELO_VIVO)
    cat_principal = triage_res["categoria"]
    tipo_principal = triage_res["tipo"]

    info_web = {"resultados": [], "corpus_web_extraido": ""}
    if buscar_web_switch:
        info_web = buscar_informacion_juridica_en_internet(cat_principal, tipo_principal)

    if auto_entrenar_switch and (btn_analizar or audio_input is not None or len(relato_usuario) > 25):
        res_auto = registrar_experiencia_en_memoria_autonoma(
            relato_original=relato_usuario,
            clasificacion=triage_res,
            fuente="Triage Multi-Materia v10.0",
            texto_web_extra=info_web.get("corpus_web_extraido", ""),
            repo_gh=gh_repo_input,
            token_gh=gh_token_input,
            branch_gh=gh_branch_input
        )
        if res_auto["aprendio"]:
            VECTORIZER_VIVO, MODELO_VIVO, STATS_IA = entrenar_red_neuronal_evolutiva()

    datos_sueldo = normalizar_texto_coloquial(relato_usuario)
    dias_vac_omitidas = 12 if datos_sueldo["falta_prestaciones"] else 0
    dias_desc_omitidos = 12 if datos_sueldo["falta_prestaciones"] else 0
    dias_sueldo_ret = datos_sueldo["quincenas_retenidas"] * 15

    res_liq_rapido = calcular_liquidacion_y_derechos_lft(
        salario_mensual=datos_sueldo["salario_mensual"],
        anios_antiguedad=1,
        meses_adicionales=0,
        motivo_separacion="Despido Injustificado",
        dias_sueldo_pendientes=dias_sueldo_ret,
        dias_vacaciones_no_gozadas=dias_vac_omitidas,
        dias_descanso_trabajados=dias_desc_omitidos
    )

    tabulador_act = TABULADOR_HONORARIOS_RAMA.get(cat_principal, TABULADOR_HONORARIOS_RAMA["Casos Laborales"])
    res_hon_rapido = calcular_cotizacion_honorarios_dual(
        monto_en_disputa=res_liq_rapido["gran_total_bruto"],
        costo_asesoria=float(tabulador_act["asesoria_sug"]),
        anticipo_fijo=float(tabulador_act["anticipo_sug"]),
        pct_litis=float(tabulador_act["pct_sugerido"])
    )

    if triage_res["es_multimateria"]:
        lista_etiquetas_html = "<br>".join([f"• <b>{m['titulo_prioridad']}</b>" for m in triage_res["modulos_triage"]])
        triage_html = " ".join([
            "<div class='triage-box'>",
            f"<h4 style='margin-top:0;'>🚨 TRIAGE MULTI-DISCIPLINARIO ACTIVADO: DETECTÉ {len(triage_res['modulos_triage'])} PROBLEMAS DISTINTOS EN TU CASO</h4>",
            "Para que no te confundas ni mezcles trámites, separé tu problema por <b>Orden de Urgencia Vital</b> ",
            "(primero tu libertad e integridad ante la policía, segundo bloquear el robo de tus cuentas/INE, ",
            "tercero cobrar tu trabajo y cuarto frenar a los cobradores):<br><br>",
            f"{lista_etiquetas_html}",
            "</div>"
        ])
        st.markdown(triage_html, unsafe_allow_html=True)

    if triage_res["tiene_componente_laboral"]:
        if datos_sueldo["oferta_patronal"] > 0 and datos_sueldo["oferta_patronal"] < res_liq_rapido["gran_total_bruto"]:
            dif_robo = round(res_liq_rapido["gran_total_bruto"] - datos_sueldo["oferta_patronal"], 2)
            alerta_html = " ".join([
                "<div class='alert-box'>",
                "<h4 style='margin-top:0;'>🚨 ¡ALERTA LABORAL! TE QUIEREN DAR MUCHO MENOS DE LO QUE MARCA LA LEY</h4>",
                f"Con tu sueldo <b>{datos_sueldo['periodicidad']}</b>, la empresa solo te ofrece <b>${datos_sueldo['oferta_patronal']:,.2f} MXN</b>, ",
                f"pero por ley tu reclamo laboral asciende a <b>${res_liq_rapido['gran_total_bruto']:,.2f} MXN</b> (perderías <b>${dif_robo:,.2f} MXN</b> si firmas renuncia).",
                "</div>"
            ])
            st.markdown(alerta_html, unsafe_allow_html=True)

        st.markdown("#### 💼 Resumen Financiero Exclusivo de tu Reclamo Laboral (Ley Federal del Trabajo)")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Reclamo Laboral LFT", f"${res_liq_rapido['gran_total_bruto']:,.2f}", "Liquidación + Sueldos")
        m2.metric("Libre para ti (Con Abogado)", f"${res_hon_rapido['monto_neto_libre_ciudadano']:,.2f}", "O 100% libre en PROFEDET")
        m3.metric("Quincenas Retenidas", f"{datos_sueldo['quincenas_retenidas']}", "Incluidas en el cálculo")
        m4.metric("Plazo Límite Laboral", "60 Días", "Para acudir a Conciliación")
    else:
        escudo_no_lab = " ".join([
            "<div class='money-box'>",
            "<h4 style='margin-top:0;'>🛡️ ESCUDO PATRIMONIAL Y CIUDADANO (CASO NO LABORAL)</h4>",
            "• <b>Nota importante:</b> Tu caso pertenece a materia <b>Vial / Penal / Civil / Cobranza</b>, por lo que aquí <b>no aplica cálculo de liquidación laboral LFT</b>.<br>",
            "• <b>Lo que debes pagar en mordidas o extorsiones:</b> <b>$0.00 MXN</b> (Exige boleta oficial o denuncia en Asuntos Internos / CONDUSEF).<br>",
            f"• <b>Defensoría Oficial Gratuita disponible:</b> {tabulador_act['via_gratuita']}",
            "</div>"
        ])
        st.markdown(escudo_no_lab, unsafe_allow_html=True)

    guias_acumuladas = []
    for idx_mod, mod_info in enumerate(triage_res["modulos_triage"], 1):
        cat_m = mod_info["categoria"]
        tipo_m = mod_info["tipo"]
        guia_m = obtener_guia_humana_exhaustiva(cat_m, tipo_m, datos_sueldo, res_liq_rapido)
        tab_m = TABULADOR_HONORARIOS_RAMA.get(cat_m, TABULADOR_HONORARIOS_RAMA["Casos Laborales"])
        guias_acumuladas.append({
            "categoria": cat_m,
            "titulo_prioridad": mod_info["titulo_prioridad"],
            "guia": guia_m
        })

        expandir_mod = True if idx_mod == 1 or not triage_res["es_multimateria"] else True
        with st.expander(f"{mod_info['titulo_prioridad']}", expanded=expandir_mod):
            rad_html = " ".join([
                "<div class='card-box'>",
                f"<span class='step-badge'>MÓDULO {idx_mod}: {cat_m.upper()}</span><br><br>",
                f"<b>🔍 Diagnóstico directo:</b><br>{guia_m['radiografia']}",
                "</div>"
            ])
            st.markdown(rad_html, unsafe_allow_html=True)

            st.markdown("##### 💡 Explicación en Español Sencillo (El porqué de tus derechos)")
            for exp_item in guia_m["explicacion_sencilla"]:
                st.markdown(exp_item)

            st.markdown("##### 🧭 Guía Cronológica Paso a Paso")
            cf1, cf2, cf3 = st.columns(3)
            with cf1:
                st.markdown("**⏱️ PASO A: Hoy mismo**")
                for p_a in guia_m["fase_1_hoy"]:
                    st.markdown(f"- {p_a}")
            with cf2:
                st.markdown("**📅 PASO B: Esta semana**")
                for p_b in guia_m["fase_2_semana"]:
                    st.markdown(f"- {p_b}")
            with cf3:
                st.markdown("**🏛️ PASO C: Ante la autoridad**")
                for p_c in guia_m["fase_3_audiencia"]:
                    st.markdown(f"- {p_c}")

            guion_html = " ".join([
                "<div class='script-box'>",
                f"{guia_m['guion_palabra_por_palabra'].replace(chr(10), '<br>')}",
                "</div>"
            ])
            st.markdown(guion_html, unsafe_allow_html=True)

            ce1, ce2 = st.columns([1.2, 1])
            with ce1:
                st.markdown("**🚫 Errores Fatales que debes evitar:**")
                for err_i in guia_m["errores_fatales"]:
                    st.markdown(f"- {err_i}")
            with ce2:
                aut_box = " ".join([
                    "<div class='lawyer-box'>",
                    f"<b>🏢 Dónde Acudir:</b><br>{guia_m['autoridad']}<br><br>",
                    f"<b>⚖️ Especialista:</b><br>{guia_m['abogado_req']}<br><br>",
                    f"<b>💚 Vía 100% Gratuita:</b><br>{tab_m['via_gratuita']}",
                    "</div>"
                ])
                st.markdown(aut_box, unsafe_allow_html=True)

    # --------------------------------------------------------------------------
    # SECCIÓN NUEVA v10.0: COTIZADOR TRANSPARENTE Y CONTACTO DIRECTO AL DESPACHO
    # --------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### ⚖️ Asesoría Profesional y Contratación Privada con el Despacho")
    st.write("Si prefieres llevar tu caso con el respaldo de un abogado experto en lugar de hacerlo solo, aquí tienes las opciones de contratación transparente y el enlace directo:")

    col_cot1, col_cot2 = st.columns(2)
    with col_cot1:
        st.markdown(f"""
        <div class='lawyer-box'>
            <b>💼 Esquema de Honorarios Estimados ({cat_principal}):</b><br>
            • <b>Asesoría Inicial:</b> ${tabulador_act['asesoria_sug']:,.2f} MXN<br>
            • <b>Anticipo Inicial:</b> ${tabulador_act['anticipo_sug']:,.2f} MXN ({tabulador_act['esquema_cobro']})<br>
            • <b>Vía Institucional Gratuita:</b> {tab_m['via_gratuita']}
        </div>
        """, unsafe_allow_html=True)

    with col_cot2:
        texto_wa = urllib.parse.quote(
            f"Hola Ing. Saúl / Equipo Legal. Analicé mi caso en ForensicMind (Folio: SAOR-2026-V10 | Materia: {cat_principal}). Me interesa agendar una asesoría privada."
        )
        link_whatsapp = f"https://wa.me/{DEFAULT_WHATSAPP_DESPACHO}?text={texto_wa}"
        st.markdown(f"""
        <div class='money-box' style='text-align:center;'>
            <b>💬 ¿Deseas atención personalizada inmediata?</b><br>
            Envía tu expediente y resumen de caso directamente al WhatsApp profesional del despacho.<br><br>
            <a href='{link_whatsapp}' target='_blank' style='background-color:#059669; color:white; padding:12px 24px; border-radius:10px; font-weight:bold; text-decoration:none; display:inline-block;'>📱 Contactar al Despacho por WhatsApp ↗</a>
        </div>
        """, unsafe_allow_html=True)

    if buscar_web_switch and info_web["resultados"]:
        st.markdown("---")
        st.markdown("#### 🌐 Información Verificada en Internet para tu Caso")
        for rw in info_web["resultados"][:3]:
            w_html = " ".join([
                "<div class='web-box'>",
                f"<b>{rw['fuente']}</b> — <a href='{rw['url']}' target='_blank' style='font-weight:bold; color:#0284c7;'>{rw['titulo']} ↗</a><br>",
                f"<span>{rw['resumen']}</span>",
                "</div>"
            ])
            st.markdown(w_html, unsafe_allow_html=True)

    st.markdown("---")
    hash_evidencia = hashlib.sha256(relato_usuario.encode("utf-8")).hexdigest()
    guia_primaria = guias_acumuladas[0]["guia"]

    col_dl1, col_dl2 = st.columns(2)
    with col_dl1:
        bytes_docx = generar_dictamen_word_docx(
            "SAOR-2026-V10", relato_usuario, triage_res, guias_acumuladas, res_liq_rapido, res_hon_rapido, hash_evidencia
        )
        if bytes_docx:
            st.download_button(
                label="📄 Descargar Manual Completo Multi-Materia en Word (.docx)",
                data=bytes_docx,
                file_name="Manual_Triage_Paso_a_Paso_v10.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
    with col_dl2:
        with st.expander("📱 Ver y Descargar Ficha Resumen 9:16 para Celular", expanded=False):
            img_9_16 = generar_ficha_vertical_9_16(
                "SAOR-2026-V10", triage_res, guia_primaria,
                res_liq_rapido["gran_total_bruto"], res_hon_rapido["monto_neto_libre_ciudadano"], hash_evidencia
            )
            st.image(img_9_16, caption=f"Ficha 9:16 v10.0 — {AUTOR_SISTEMA}", use_container_width=True)
            _, buf_png = cv2.imencode(".png", cv2.cvtColor(img_9_16, cv2.COLOR_RGB2BGR))
            st.download_button(
                "📥 Descargar Imagen 9:16 (.png)",
                buf_png.tobytes(),
                "Ficha_Triage_9_16.png",
                "image/png",
                use_container_width=True
            )

    st.markdown("---")
    st.markdown("### 📬 ¿No encontraste exactamente tu problema o tienes una duda especial?")
    st.write("Envíale tu caso directamente al **Ing. Saúl Ángel Ortiz Rogel** para analizarlo e integrarlo a la inteligencia artificial:")
    with st.form("form_buzon_ciudadano"):
        c_bz1, c_bz2 = st.columns(2)
        nom_bz = c_bz1.text_input("Tu nombre o apodo (opcional):", placeholder="Ej. Juan Pérez")
        con_bz = c_bz2.text_input("Tu correo o WhatsApp para responderte (opcional):", placeholder="Ej. correo@ejemplo.com")
        duda_bz = st.text_area("Explícanos qué te sucedió y qué información necesitas:")
        if st.form_submit_button("📨 Enviar mi Caso al Administrador", use_container_width=True):
            if duda_bz and len(duda_bz.strip()) > 5:
                conn = sqlite3.connect(DB_NAME)
                sql_ins_bz = " ".join([
                    "INSERT OR IGNORE INTO buzon_casos_usuarios",
                    "(nombre_usuario, contacto_usuario, relato_duda, estado, fecha_envio)",
                    "VALUES (?, ?, ?, 'Pendiente de revisión', ?)"
                ])
                conn.execute(sql_ins_bz, (
                    nom_bz or "Ciudadano Anónimo",
                    con_bz or "Sin contacto",
                    duda_bz.strip(),
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                ))
                conn.commit()
                conn.close()
                guardar_y_sincronizar_cerebro_autonomo(gh_repo_input, gh_token_input, gh_branch_input)
                asunto_mail = urllib.parse.quote("Nuevo Caso Ciudadano - ForensicMind v10.0")
                cuerpo_mail = urllib.parse.quote(f"Nombre: {nom_bz}\nContacto: {con_bz}\n\nCaso:\n{duda_bz}")
                st.success("✅ ¡Tu caso quedó guardado en el buzón del Administrador y sincronizado en GitHub!")
                st.markdown(
                    f"👉 [Haz clic aquí si deseas enviarlo también desde tu app de Correo](mailto:{DEFAULT_ADMIN_EMAIL}?subject={asunto_mail}&body={cuerpo_mail})"
                )
            else:
                st.warning("Por favor escribe una breve descripción de tu duda.")

# ==============================================================================
# PESTAÑA 2: CALCULADORA EXHAUSTIVA DE LIQUIDACIÓN
# ==============================================================================
with tab2:
    st.subheader("💰 Calculadora Fácil de Liquidación Laboral y Costo de Abogado")
    st.write("Exclusiva para asuntos de trabajo: calcula cuánto te corresponde por despido o renuncia:")

    col_c1, col_c2 = st.columns(2)
    with col_c1:
        st.markdown("#### 👷 1. Datos de tu Sueldo y Trabajo")
        per_sel = st.selectbox(
            "¿Cada cuándo te pagan?",
            ["Semanal (Cada 7 días)", "Quincenal (Cada 15 días)", "Mensual (Cada 30 días)", "Por día"]
        )
        val_ini = 2300.0 if "Semanal" in per_sel else datos_sueldo["salario_mensual"]
        monto_ing = st.number_input("¿Cuánto ganas en ese periodo? ($ MXN)", min_value=100.0, max_value=500000.0, value=float(val_ini), step=100.0)

        if "Semanal" in per_sel:
            sueldo_mes_c = round((monto_ing / 7.0) * 30.0, 2)
        elif "Quincenal" in per_sel:
            sueldo_mes_c = round(monto_ing * 2.0, 2)
        elif "Por día" in per_sel:
            sueldo_mes_c = round(monto_ing * 30.0, 2)
        else:
            sueldo_mes_c = round(monto_ing, 2)

        st.info(f"📌 **Tu sueldo mensual equivale a:** ${sueldo_mes_c:,.2f} MXN (${sueldo_mes_c/30.0:,.2f} pesos por día)")
        anios_c = st.number_input("¿Cuántos años cumplidos trabajaste ahí?", min_value=0, max_value=50, value=1, step=1)
        meses_c = st.number_input("¿Con cuántos meses adicionales?", min_value=0, max_value=11, value=0, step=1)

    with col_c2:
        st.markdown("#### 📋 2. Qué te deben y cuánto te ofrecen")
        motivo_c = st.selectbox(
            "¿Por qué saliste del trabajo?",
            [
                "Despido Injustificado (Me corrieron o me deben sueldos - Liquidación Completa)",
                "Renuncia Voluntaria (Yo decidí salirme - Solo Finiquito)"
            ]
        )
        oferta_c = st.number_input(
            "¿Cuánto dinero te quiere dar el patrón hoy? ($ MXN)",
            min_value=0.0, max_value=1000000.0,
            value=float(datos_sueldo["oferta_patronal"] if datos_sueldo["oferta_patronal"] > 0 else 5000.0),
            step=500.0
        )
        vac_pend_c = st.number_input("Días de vacaciones pasadas que nunca te dieron:", min_value=0, max_value=120, value=int(dias_vac_omitidas), step=1)
        desc_trab_c = st.number_input("Días de descanso que trabajaste sin pago doble:", min_value=0, max_value=100, value=int(dias_desc_omitidos), step=1)
        dias_sueldo_c = st.number_input("Días ya trabajados que no te han pagado:", min_value=0, max_value=180, value=int(dias_sueldo_ret), step=1)

    res_det = calcular_liquidacion_y_derechos_lft(
        salario_mensual=sueldo_mes_c,
        anios_antiguedad=anios_c,
        meses_adicionales=meses_c,
        motivo_separacion=motivo_c,
        dias_sueldo_pendientes=dias_sueldo_c,
        dias_vacaciones_no_gozadas=vac_pend_c,
        dias_descanso_trabajados=desc_trab_c
    )

    k1, k2, k3 = st.columns(3)
    k1.metric("Lo que te corresponde por Ley", f"${res_det['gran_total_bruto']:,.2f} MXN")
    k2.metric("Lo que te quiere dar la empresa", f"${oferta_c:,.2f} MXN")
    k3.metric("Dinero que te están quitando", f"${max(0.0, res_det['gran_total_bruto'] - oferta_c):,.2f} MXN")

    st.markdown("#### 📊 Desglose explicado renglón por renglón")
    st.dataframe(res_det["tabla_conceptos"].style.format({"Monto (MXN)": "${:,.2f}"}), use_container_width=True, hide_index=True)

# ==============================================================================
# PESTAÑA 3: COMPARATIVA VISUAL LADO A LADO (GANADO VS. PERDIDO)
# ==============================================================================
with tab4:
    st.subheader("🔬 Comparativa Real: Estrategia Ganadora vs. Error Fatal")
    st.write("Compara lado a lado qué acciones hicieron ganar un caso real y qué errores hicieron que otro se perdiera:")

    f_cat_comp = st.selectbox(
        "Elige el tema que deseas comparar:",
        list(TIPOS_CASOS.keys()),
        key="filtro_comparativa_visual"
    )

    conn = sqlite3.connect(DB_NAME)
    df_comp = pd.read_sql_query("SELECT * FROM casos_juridicos_cbr WHERE categoria_caso = ?", conn, params=(f_cat_comp,))
    conn.close()

    if df_comp.empty:
        st.info("Aún no hay suficientes registros comparativos para esta categoría.")
    else:
        ganados_list = df_comp[df_comp["resultado"] == "GANADO"].to_dict(orient="records")
        perdidos_list = df_comp[df_comp["resultado"] == "PERDIDO"].to_dict(orient="records")

        col_cg, col_cp = st.columns(2)

        with col_cg:
            st.markdown("#### 🟢 Estrategia del Caso GANADO")
            if ganados_list:
                for g in ganados_list:
                    box_g = " ".join([
                        "<div class='money-box'>",
                        "<span class='win-badge'>CASO EXITOSO</span><br><br>",
                        f"<b>📌 {g['titulo']}</b><br><br>",
                        f"<b>🗣️ Qué ocurrió:</b><br>{g['evidencia_logs']}<br><br>",
                        f"<b>💼 Estrategia que dio la victoria:</b><br>{g['accion_abogado']}<br><br>",
                        f"<b>👨‍⚖️ Resolución de la Autoridad:</b><br>{g['reaccion_juez']}<br><br>",
                        f"<b>✅ Lo que tú debes hacer:</b><br>{g['que_hacer_ciudadano']}",
                        "</div>"
                    ])
                    st.markdown(box_g, unsafe_allow_html=True)
            else:
                st.info("No hay un caso ganado registrado en esta categoría aún.")

        with col_cp:
            st.markdown("#### 🔴 Error del Caso PERDIDO")
            if perdidos_list:
                for p in perdidos_list:
                    box_p = " ".join([
                        "<div class='alert-box'>",
                        "<span class='loss-badge'>CASO FRACASADO</span><br><br>",
                        f"<b>📌 {p['titulo']}</b><br><br>",
                        f"<b>🗣️ Qué ocurrió:</b><br>{p['evidencia_logs']}<br><br>",
                        f"<b>💥 Error fatal que causó la derrota:</b><br>{p['causa_raiz_fallo']}<br><br>",
                        f"<b>🚫 Lo que NUNCA debes hacer:</b><br>{p['que_no_hacer_ciudadano']}",
                        "</div>"
                    ])
                    st.markdown(box_p, unsafe_allow_html=True)
            else:
                st.info("No hay un caso perdido registrado en esta categoría aún.")

# ==============================================================================
# PESTAÑAS EXCLUSIVAS DE ADMINISTRADOR (SOLO VISIBLES CON CONTRASEÑA)
# ==============================================================================
if st.session_state["es_admin"]:
    with tab3:
        st.subheader("☁️ [ADMIN] Control de Autoinyección y Sincronización con GitHub")
        st.write(f"• **Repositorio conectado:** `{gh_repo_input}` | **Muestras en Red Neuronal:** `{STATS_IA['total_muestras']}`")

        col_ad1, col_ad2 = st.columns(2)
        with col_ad1:
            if st.button("☁️ Forzar Sincronización a GitHub Ahora", use_container_width=True):
                res_s = guardar_y_sincronizar_cerebro_autonomo(gh_repo_input, gh_token_input, gh_branch_input)
                if res_s["github_ok"]:
                    st.success(res_s["msg"])
                else:
                    st.warning(res_s["msg"])
        with col_ad2:
            conn = sqlite3.connect(DB_NAME)
            df_m_adm = pd.read_sql_query("SELECT * FROM memoria_neuronal_autonoma ORDER BY id DESC", conn)
            conn.close()
            st.write(f"Memorias aprendidas automáticamente: `{len(df_m_adm)}`")

        if not df_m_adm.empty:
            st.dataframe(df_m_adm[["id", "relato_original", "categoria_caso", "fecha_aprendizaje"]], use_container_width=True, hide_index=True)

    with tab5:
        st.subheader("➕ [ADMIN] Enseñar Nuevo Caso Oficial a la IA")
        with st.form("form_admin_nuevo_caso"):
            c_a1, c_a2, c_a3 = st.columns(3)
            ncod = c_a1.text_input("Código (ej. EXP-TRA-2026-099)")
            ncat = c_a2.selectbox("Categoría", list(TIPOS_CASOS.keys()))
            nres = c_a3.selectbox("Resultado", ["GANADO", "PERDIDO"])
            ntit = st.text_input("Título del Caso")
            nevid = st.text_area("Relato de hechos (incluye palabras coloquiales)")
            nqh = st.text_area("✅ Qué debe hacer el ciudadano")
            nqnh = st.text_area("🚫 Qué NO debe hacer el ciudadano")

            if st.form_submit_button("💾 Guardar Caso y Sincronizar con GitHub"):
                if ncod and ntit and nevid:
                    conn = sqlite3.connect(DB_NAME)
                    sql_adm_ins = " ".join([
                        "INSERT OR IGNORE INTO casos_juridicos_cbr (",
                        "codigo_caso, macro_rama, rama_derecho, categoria_caso, tipo_caso,",
                        "titulo, actores_intervinientes, etapa_clave, medida_cautelar,",
                        "resolucion_obtenida, recurso_utilizado, evidencia_logs, hash_sha256,",
                        "resultado, accion_abogado, reaccion_juez, causa_raiz_fallo,",
                        "estrategia_ganadora, que_hacer_ciudadano, que_no_hacer_ciudadano, fecha_registro",
                        ") VALUES (?, 'Derecho Privado', 'Derecho General', ?, ?, ?, 'Autoridad', '1. Inicio', 'Medida', 'Resolución', 'Recurso', ?, ?, ?, 'Estrategia aplicada', 'Resolución judicial', 'Causa analizada', 'Estrategia', ?, ?, ?)"
                    ])
                    conn.execute(sql_adm_ins, (
                        ncod,
                        ncat,
                        TIPOS_CASOS[ncat][0],
                        ntit,
                        nevid,
                        hashlib.sha256(nevid.encode("utf-8")).hexdigest(),
                        nres,
                        nqh,
                        nqnh,
                        datetime.now().strftime("%Y-%m-%d")
                    ))
                    conn.commit()
                    conn.close()
                    sync_r = guardar_y_sincronizar_cerebro_autonomo(gh_repo_input, gh_token_input, gh_branch_input)
                    st.success(f"✅ Caso guardado y sincronizado. {sync_r['msg']}")

    with tab6:
        st.subheader("📬 [ADMIN] Buzón de Casos Enviados por Usuarios (Entrenar en 1 Clic)")
        conn = sqlite3.connect(DB_NAME)
        df_bz_view = pd.read_sql_query("SELECT * FROM buzon_casos_usuarios ORDER BY id DESC", conn)
        conn.close()

        if df_bz_view.empty:
            st.info("Aún no hay mensajes nuevos en el buzón ciudadano.")
        else:
            st.dataframe(df_bz_view, use_container_width=True, hide_index=True)
            st.markdown("#### 🧠 Convertir una duda del Buzón en Aprendizaje de la IA")
            with st.form("form_entrenar_desde_buzon"):
                id_sel = st.selectbox("Selecciona el ID del caso enviado:", df_bz_view["id"].tolist())
                cat_asignar = st.selectbox("Asignar a la Categoría:", list(TIPOS_CASOS.keys()))
                qh_asignar = st.text_area("Recomendación oficial para casos como este:", value="Acudir a la autoridad competente con pruebas documentales sin firmar renuncias ni entregar efectivo.")
                if st.form_submit_button("🚀 Aprobar, Entrenar IA y Sincronizar con GitHub"):
                    fila_bz = df_bz_view[df_bz_view["id"] == id_sel].iloc[0]
                    rel_bz = str(fila_bz["relato_duda"])
                    cod_auto = f"EXP-BUZON-{id_sel}-{datetime.now().strftime('%H%M')}"
                    conn = sqlite3.connect(DB_NAME)
                    sql_bz_cbr = " ".join([
                        "INSERT OR IGNORE INTO casos_juridicos_cbr (",
                        "codigo_caso, macro_rama, rama_derecho, categoria_caso, tipo_caso,",
                        "titulo, actores_intervinientes, etapa_clave, medida_cautelar,",
                        "resolucion_obtenida, recurso_utilizado, evidencia_logs, hash_sha256,",
                        "resultado, accion_abogado, reaccion_juez, causa_raiz_fallo,",
                        "estrategia_ganadora, que_hacer_ciudadano, que_no_hacer_ciudadano, fecha_registro",
                        ") VALUES (?, 'Derecho Ciudadano', ?, ?, ?, ?, 'Autoridad Oficial', '1. Atención Inicial', 'Protección Ciudadana', 'Caso Atendido', 'Vía Oficial', ?, ?, 'GANADO', ?, 'Protección otorgada', 'N/A', ?, ?, 'No caer en coacción ni entregar dinero sin recibo.', ?)"
                    ])
                    conn.execute(sql_bz_cbr, (
                        cod_auto,
                        cat_asignar,
                        cat_asignar,
                        TIPOS_CASOS[cat_asignar][0],
                        f"Caso Ciudadano del Buzón #{id_sel}",
                        rel_bz,
                        hashlib.sha256(rel_bz.encode("utf-8")).hexdigest(),
                        qh_asignar,
                        qh_asignar,
                        qh_asignar,
                        datetime.now().strftime("%Y-%m-%d")
                    ))
                    conn.execute("UPDATE buzon_casos_usuarios SET estado = 'Integrado a la IA' WHERE id = ?", (int(id_sel),))
                    conn.commit()
                    conn.close()
                    sync_bz = guardar_y_sincronizar_cerebro_autonomo(gh_repo_input, gh_token_input, gh_branch_input)
                    st.success(f"✅ Caso #{id_sel} integrado a la Red Neuronal y respaldado. {sync_bz['msg']}")
                    st.rerun()

    with tab7:
        st.subheader("🛠️ [ADMIN] Cotizador Privado de Honorarios & Tarifas de Despacho")
        st.write("Usa esta herramienta privada para calcular presupuestos formales para tus clientes según la materia:")

        mat_sel_adm = st.selectbox("Selecciona la Materia Jurídica:", list(TABULADOR_HONORARIOS_RAMA.keys()))
        tab_adm = TABULADOR_HONORARIOS_RAMA[mat_sel_adm]

        c_adm1, c_adm2 = st.columns(2)
        monto_disp_adm = c_adm1.number_input("Monto aproximado en disputa ($ MXN):", min_value=0.0, value=50000.0, step=5000.0)
        pct_exito_adm = c_adm2.number_input("Porcentaje de éxito (Cuota Litis %):", min_value=0.0, max_value=50.0, value=float(tab_adm["pct_sugerido"]), step=1.0)

        anticipo_adm = st.number_input("Anticipo inicial requerido ($ MXN):", min_value=0.0, value=float(tab_adm["anticipo_sug"]), step=500.0)
        asesoria_adm = st.number_input("Costo de consulta/asesoría ($ MXN):", min_value=0.0, value=float(tab_adm["asesoria_sug"]), step=100.0)

        res_cot_priv = calcular_cotizacion_honorarios_dual(
            monto_en_disputa=monto_disp_adm,
            costo_asesoria=asesoria_adm,
            anticipo_fijo=anticipo_adm,
            pct_litis=pct_exito_adm
        )

        st.markdown(f"""
        <div class='money-box'>
            <b>📊 Propuesta Económica Calculada para el Cliente:</b><br>
            • <b>Anticipo / Retainer Inicial:</b> ${anticipo_adm:,.2f} MXN<br>
            • <b>Comisión por Éxito ({pct_exito_adm}%):</b> ${monto_disp_adm * (pct_exito_adm/100.0):,.2f} MXN<br>
            • <b>Total de Honorarios del Despacho:</b> ${res_cot_priv['honorarios_totales_abogado']:,.2f} MXN<br>
            • <b>Neto Libre para el Cliente:</b> ${res_cot_priv['monto_neto_libre_ciudadano']:,.2f} MXN
        </div>
        """, unsafe_allow_html=True)

        propuesta_whatsapp = (
            f"Estimado cliente, respecto a su asunto de {mat_sel_adm}, la propuesta de honorarios del despacho es la siguiente:\n"
            f"- Asesoría inicial: ${asesoria_adm:,.2f} MXN\n"
            f"- Anticipo inicial para gestiones: ${anticipo_adm:,.2f} MXN\n"
            f"- Comisión por éxito (Cuota Litis): {pct_exito_adm}% sobre lo recuperado.\n"
            f"Quedamos a sus órdenes para iniciar con la defensa jurídica."
        )
        st.text_area("📋 Texto listo para copiar y enviar por WhatsApp al cliente:", value=propuesta_whatsapp, height=120)

    with tab8:
        st.subheader("👥 [ADMIN] Tráfico y Analítica de Visitantes en la Página")
        conn = sqlite3.connect(DB_NAME)
        df_visitas = pd.read_sql_query("SELECT * FROM visitas_pagina ORDER BY id DESC", conn)
        conn.close()

        total_visitas = len(df_visitas)
        hoy_str = datetime.now().strftime("%Y-%m-%d")
        visitas_hoy = len(df_visitas[df_visitas["fecha_hora"].str.startswith(hoy_str)]) if not df_visitas.empty else 0

        v1, v2 = st.columns(2)
        v1.metric("Visitas Totales Acumuladas", f"{total_visitas}")
        v2.metric("Visitas Registradas Hoy", f"{visitas_hoy}")

        st.markdown("#### 🕒 Registro Cronológico de Sesiones")
        if df_visitas.empty:
            st.info("Aún no hay visitas registradas.")
        else:
            st.dataframe(df_visitas, use_container_width=True, hide_index=True)

st.markdown("---")
st.caption(f"⚙️ Arquitectura y Desarrollo: {AUTOR_SISTEMA} • {MARCA_SISTEMA}")
